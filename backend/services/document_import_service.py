from __future__ import annotations

import json
import os
import re
import uuid
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from werkzeug.datastructures import FileStorage
from werkzeug.utils import secure_filename

from backend.config import Config
from backend.models.database import get_connection, row_to_dict


PIPELINE_IMPORT_DOCUMENT = "import_document"
STATUS_PENDING = "pending"
STATUS_UPLOADED = "uploaded"
STATUS_PARSED = "parsed"
STATUS_CHUNKED = "chunked"
STATUS_READY_FOR_REVIEW = "ready_for_review"
STATUS_PUBLISHED = "published"
STATUS_FAILED = "failed"
VALIDITY_UNKNOWN = "unknown"
REQUIRED_PUBLISH_FIELDS = (
    "title",
    "document_number",
    "document_type",
    "issuing_body",
    "issued_date",
    "effective_date",
    "validity_status",
)
METADATA_UPDATE_FIELDS = {
    "document_number",
    "title",
    "source_url",
    "issuing_body",
    "signer_title",
    "signer_name",
    "document_type",
    "issued_date",
    "effective_date",
    "expiry_date",
    "validity_status",
}


class DocumentImportError(Exception):
    def __init__(
        self,
        code: str,
        message: str,
        status_code: int = 400,
        details: dict[str, Any] | None = None,
    ):
        super().__init__(message)
        self.code = code
        self.message = message
        self.status_code = status_code
        self.details = details or {}


@dataclass(frozen=True)
class ImportMetadata:
    document_number: str
    title: str
    source_url: str | None
    issuing_body: str | None
    issued_date: str | None
    effective_date: str | None
    expiry_date: str | None
    validity_status: str
    signer_title: str | None
    signer_name: str | None
    document_type: str | None
    relations: list[dict[str, Any]]
    relations_unavailable: bool
    needs_review: bool
    needs_review_fields: list[str]
    extraction: dict[str, Any]


class DocumentImportService:
    def __init__(self, config: Config):
        self.config = config
        self.db_path = config.sqlite_db_path

    def import_document(
        self,
        file_storage: FileStorage | None,
        form_data: dict[str, Any],
        requested_by_user_id: str,
    ) -> dict[str, Any]:
        pipeline_run_id = str(uuid.uuid4())
        import_batch_id = str(uuid.uuid4())
        now = utc_now_iso()

        self._create_pipeline_run(
            pipeline_run_id=pipeline_run_id,
            requested_by_user_id=requested_by_user_id,
            status=STATUS_PENDING,
            input_json={
                "document_number": optional_str(form_data.get("document_number")),
                "title": optional_str(form_data.get("title")),
                "issued_date": optional_str(form_data.get("issued_date")),
                "effective_date": optional_str(form_data.get("effective_date")),
                "expiry_date": optional_str(form_data.get("expiry_date")),
                "validity_status": optional_str(form_data.get("validity_status"))
                or VALIDITY_UNKNOWN,
                "import_batch_id": import_batch_id,
                "relations_json_provided": bool(
                    optional_str(form_data.get("relations_json"))
                ),
            },
            now=now,
        )

        try:
            self._validate_docx(file_storage)
            assert file_storage is not None

            raw_docx_path = self._save_raw_docx(
                file_storage=file_storage,
                import_batch_id=import_batch_id,
            )
            self._mark_pipeline(
                pipeline_run_id,
                STATUS_UPLOADED,
                "Raw DOCX uploaded.",
                {"raw_docx_path": str(raw_docx_path)},
            )

            try:
                extracted = extract_docx(raw_docx_path)
            except Exception as exc:  # pragma: no cover - exact parser errors vary.
                raise DocumentImportError(
                    "DOCX_PARSE_FAILED",
                    "Cannot parse uploaded DOCX.",
                    details={"reason": str(exc)},
                ) from exc

            text = "\n".join(extracted["paragraphs"])
            if not text.strip():
                raise DocumentImportError(
                    "DOCX_PARSE_FAILED",
                    "Uploaded DOCX does not contain extractable paragraph text.",
                )

            metadata_hints = extract_metadata_hints(
                text=text,
                tables=extracted["tables"],
                paragraphs=extracted["paragraphs"],
            )
            metadata = self._parse_metadata(
                form_data=form_data,
                metadata_hints=metadata_hints,
                text=text,
                paragraphs=extracted["paragraphs"],
                import_batch_id=import_batch_id,
            )
            document_id, version = self._upsert_document_registry(metadata, now)
            preprocessed_text_path = self._write_preprocessed_text(
                document_id=document_id,
                import_batch_id=import_batch_id,
                text=text,
            )
            merged_metadata = build_normalized_metadata(
                document_id=document_id,
                version=version,
                import_batch_id=import_batch_id,
                metadata=metadata,
                metadata_hints=metadata_hints,
            )
            self._mark_pipeline(
                pipeline_run_id,
                STATUS_PARSED,
                "DOCX parsed and metadata hints extracted.",
                {
                    "preprocessed_text_path": str(preprocessed_text_path),
                    "tables_count": len(extracted["tables"]),
                    "metadata_hints": metadata_hints,
                    "resolved_document_number": metadata.document_number,
                    "resolved_title": metadata.title,
                    "needs_review": metadata.needs_review,
                    "needs_review_fields": metadata.needs_review_fields,
                },
            )

            warnings: list[str] = []
            chunks = self._chunk_text(
                text=text,
                document_id=document_id,
                import_batch_id=import_batch_id,
                metadata=merged_metadata,
                warnings=warnings,
            )
            validate_chunks(chunks)
            chunk_json_path = self._write_chunks(
                document_id=document_id,
                import_batch_id=import_batch_id,
                chunks=chunks,
            )

            if metadata.relations_unavailable:
                warnings.append("Document-level relations are not available.")
            if metadata.needs_review:
                warnings.append(
                    "Metadata was inferred with missing or low-confidence fields; admin review is required."
                )

            self._store_version_and_relations(
                document_id=document_id,
                version=version,
                import_batch_id=import_batch_id,
                raw_docx_path=raw_docx_path,
                preprocessed_text_path=preprocessed_text_path,
                chunk_json_path=chunk_json_path,
                metadata_json={
                    **merged_metadata,
                    "metadata_hints": metadata_hints,
                    "warnings": warnings,
                    "relations_unavailable": metadata.relations_unavailable,
                    "needs_review": metadata.needs_review,
                    "needs_review_fields": metadata.needs_review_fields,
                    "metadata_extraction": metadata.extraction,
                },
                relations=metadata.relations,
                now=utc_now_iso(),
            )
            self._mark_pipeline(
                pipeline_run_id,
                STATUS_CHUNKED,
                "DOCX chunked into article/clause records.",
                {
                    "chunk_json_path": str(chunk_json_path),
                    "chunks_count": len(chunks),
                    "warnings": warnings,
                },
            )
            self._mark_pipeline(
                pipeline_run_id,
                STATUS_READY_FOR_REVIEW,
                "Import is ready for admin review.",
                {
                    "document_id": document_id,
                    "version": version,
                    "import_batch_id": import_batch_id,
                },
            )
            publish_blockers = get_publish_blockers(
                {
                    "title": metadata.title,
                    "document_number": metadata.document_number,
                    "document_type": metadata.document_type,
                    "issuing_body": metadata.issuing_body,
                    "issued_date": metadata.issued_date,
                    "effective_date": metadata.effective_date,
                    "validity_status": metadata.validity_status,
                }
            )

            return {
                "document_id": document_id,
                "version": version,
                "import_batch_id": import_batch_id,
                "pipeline_run_id": pipeline_run_id,
                "status": STATUS_READY_FOR_REVIEW,
                "publish_blockers": publish_blockers,
                "is_publishable": not publish_blockers,
                "document_number": metadata.document_number,
                "title": metadata.title,
                "raw_docx_path": str(raw_docx_path),
                "preprocessed_text_path": str(preprocessed_text_path),
                "chunk_json_path": str(chunk_json_path),
                "warnings": warnings,
                "relations_unavailable": metadata.relations_unavailable,
                "needs_review": metadata.needs_review,
                "needs_review_fields": metadata.needs_review_fields,
            }
        except DocumentImportError as exc:
            self._fail_pipeline(
                pipeline_run_id,
                exc.message,
                {"code": exc.code, **exc.details},
            )
            raise
        except Exception as exc:
            self._fail_pipeline(
                pipeline_run_id,
                "Unexpected import failure.",
                {"code": "IMPORT_FAILED", "reason": str(exc)},
            )
            raise

    def list_documents(self, status_filter: str | None = None) -> list[dict[str, Any]]:
        with get_connection(self.db_path) as connection:
            rows = connection.execute(
                """
                SELECT
                    dr.*,
                    dv.version AS latest_version,
                    dv.status AS latest_status,
                    dv.import_batch_id AS latest_import_batch_id,
                    dv.chunk_json_path AS latest_chunk_json_path,
                    dv.metadata_json AS latest_metadata_json,
                    av.chunk_json_path AS active_chunk_json_path
                FROM document_registry dr
                LEFT JOIN document_versions dv
                    ON dv.document_id = dr.document_id
                   AND dv.version = (
                       SELECT MAX(version)
                       FROM document_versions
                       WHERE document_id = dr.document_id
                   )
                LEFT JOIN document_versions av
                    ON av.document_id = dr.document_id
                   AND av.version = dr.active_version
                WHERE dr.is_deleted = 0
                ORDER BY dr.updated_at DESC
                """
            ).fetchall()
        documents = []
        for row in rows:
            document = dict(row)
            if status_filter == "published" and not document.get("is_published"):
                continue
            if status_filter == STATUS_READY_FOR_REVIEW and document.get("latest_status") != STATUS_READY_FOR_REVIEW:
                continue

            latest_metadata = parse_json(document.get("latest_metadata_json"))
            if status_filter == STATUS_READY_FOR_REVIEW:
                overlay_metadata(document, latest_metadata)

            metadata_json = latest_metadata or parse_json(document.get("raw_metadata_json"))
            chunk_json_path = (
                document.get("active_chunk_json_path")
                if status_filter == "published"
                else document.get("latest_chunk_json_path")
            )
            document["chunk_count"] = count_chunks(Path(chunk_json_path)) if chunk_json_path else 0
            document["needs_review"] = bool(metadata_json.get("needs_review"))
            document["needs_review_fields"] = metadata_json.get("needs_review_fields", [])
            document["needs_republish"] = bool(metadata_json.get("needs_republish"))
            document["last_publish_error"] = metadata_json.get("last_publish_error")
            document["publish_blockers"] = get_publish_blockers(
                {**document, **metadata_json}
            )
            document["is_publishable"] = not document["publish_blockers"]
            documents.append(document)
        return documents

    def get_document_detail(
        self,
        document_id: str,
        chunk_preview_limit: int = 5,
        scope: str = "latest",
    ) -> dict[str, Any] | None:
        with get_connection(self.db_path) as connection:
            registry = row_to_dict(
                connection.execute(
                    "SELECT * FROM document_registry WHERE document_id = ?",
                    (document_id,),
                ).fetchone()
            )
            if registry is None:
                return None
            if registry.get("is_deleted"):
                return None

            version_where = "ORDER BY version DESC LIMIT 1"
            params: tuple[Any, ...] = (document_id,)
            if scope == "active" and registry.get("active_version"):
                version_where = "AND version = ? LIMIT 1"
                params = (document_id, registry["active_version"])
            version = row_to_dict(
                connection.execute(
                    f"""
                    SELECT *
                    FROM document_versions
                    WHERE document_id = ?
                    {version_where}
                    """,
                    params,
                ).fetchone()
            )
            relations = [
                dict(row)
                for row in connection.execute(
                    """
                    SELECT *
                    FROM document_relations
                    WHERE source_document_id = ?
                    ORDER BY created_at DESC
                    """,
                    (document_id,),
                ).fetchall()
            ]
            events = []
            if version is not None:
                events = [
                    dict(row)
                    for row in connection.execute(
                        """
                        SELECT pe.*
                        FROM pipeline_events pe
                        JOIN pipeline_runs pr ON pr.id = pe.pipeline_run_id
                        WHERE json_extract(pr.input_json, '$.import_batch_id') = ?
                        ORDER BY pe.created_at ASC
                        """,
                        (version["import_batch_id"],),
                    ).fetchall()
                ]

        chunks_preview = []
        chunks = []
        if version and version.get("chunk_json_path"):
            chunks = load_chunks(Path(version["chunk_json_path"]))
            chunks_preview = load_chunk_preview(
                Path(version["chunk_json_path"]),
                limit=chunk_preview_limit,
            )

        metadata_json = parse_json(version.get("metadata_json") if version else None)
        document = dict(registry)
        if scope != "active":
            overlay_metadata(document, metadata_json)
        publish_blockers = get_publish_blockers({**document, **metadata_json})
        return {
            "document": document,
            "version": version,
            "relations": relations,
            "pipeline_events": events,
            "chunks": chunks,
            "chunks_preview": chunks_preview,
            "chunk_count": len(chunks),
            "needs_review": bool(metadata_json.get("needs_review")),
            "needs_review_fields": metadata_json.get("needs_review_fields", []),
            "needs_republish": bool(metadata_json.get("needs_republish")),
            "last_publish_error": metadata_json.get("last_publish_error"),
            "publish_blockers": publish_blockers,
            "is_publishable": not publish_blockers,
        }

    def list_pipeline_runs(self) -> list[dict[str, Any]]:
        with get_connection(self.db_path) as connection:
            rows = connection.execute(
                """
                SELECT *
                FROM pipeline_runs
                ORDER BY created_at DESC
                """
            ).fetchall()
        return [dict(row) for row in rows]

    def get_pipeline_detail(self, run_id: str) -> dict[str, Any] | None:
        with get_connection(self.db_path) as connection:
            run = row_to_dict(
                connection.execute(
                    "SELECT * FROM pipeline_runs WHERE id = ?",
                    (run_id,),
                ).fetchone()
            )
            if run is None:
                return None
            events = [
                dict(row)
                for row in connection.execute(
                    """
                    SELECT *
                    FROM pipeline_events
                    WHERE pipeline_run_id = ?
                    ORDER BY created_at ASC
                    """,
                    (run_id,),
                ).fetchall()
            ]
        return {"pipeline_run": run, "events": events}

    def get_raw_docx_path(self, document_id: str) -> Path | None:
        with get_connection(self.db_path) as connection:
            row = row_to_dict(
                connection.execute(
                    """
                    SELECT dv.raw_docx_path
                    FROM document_registry dr
                    JOIN document_versions dv ON dv.document_id = dr.document_id
                    WHERE dr.document_id = ? AND dr.is_deleted = 0
                    ORDER BY dv.version DESC
                    LIMIT 1
                    """,
                    (document_id,),
                ).fetchone()
            )
        if not row or not row.get("raw_docx_path"):
            return None
        path = Path(row["raw_docx_path"])
        return path.resolve() if path.exists() else None

    def mark_latest_version_needs_republish(
        self,
        document_id: str,
        message: str,
        details: dict[str, Any] | None = None,
    ) -> None:
        with get_connection(self.db_path) as connection:
            version = self._get_latest_version(connection, document_id)
            if version is None:
                return
            metadata_json = parse_json(version.get("metadata_json"))
            metadata_json["needs_republish"] = True
            metadata_json["last_publish_error"] = {
                "message": message,
                "details": details or {},
                "at": utc_now_iso(),
            }
            connection.execute(
                """
                UPDATE document_versions
                SET metadata_json = ?, status = ?
                WHERE id = ?
                """,
                (
                    json.dumps(metadata_json, ensure_ascii=False),
                    STATUS_READY_FOR_REVIEW,
                    version["id"],
                ),
            )
            connection.commit()

    def _ensure_editable_latest_version(
        self,
        connection,
        document_id: str,
    ) -> tuple[dict[str, Any], dict[str, Any]]:
        registry = row_to_dict(
            connection.execute(
                "SELECT * FROM document_registry WHERE document_id = ? AND is_deleted = 0",
                (document_id,),
            ).fetchone()
        )
        if registry is None:
            raise DocumentImportError("DOCUMENT_NOT_FOUND", "Document not found.", 404)
        latest = self._get_latest_version(connection, document_id)
        if latest is None:
            raise DocumentImportError("DOCUMENT_NOT_FOUND", "Document not found.", 404)

        if (
            registry.get("is_published")
            and int(latest.get("version") or 0) == int(registry.get("active_version") or 0)
            and latest.get("status") == STATUS_PUBLISHED
        ):
            latest = self._clone_active_version_for_edit(connection, registry, latest)
        return registry, latest

    def _get_latest_version(self, connection, document_id: str) -> dict[str, Any] | None:
        return row_to_dict(
            connection.execute(
                """
                SELECT *
                FROM document_versions
                WHERE document_id = ?
                ORDER BY version DESC
                LIMIT 1
                """,
                (document_id,),
            ).fetchone()
        )

    def _clone_active_version_for_edit(
        self,
        connection,
        registry: dict[str, Any],
        active_version: dict[str, Any],
    ) -> dict[str, Any]:
        document_id = registry["document_id"]
        now = utc_now_iso()
        next_version_row = connection.execute(
            """
            SELECT COALESCE(MAX(version), 0) + 1 AS next_version
            FROM document_versions
            WHERE document_id = ?
            """,
            (document_id,),
        ).fetchone()
        version = int(next_version_row["next_version"])
        import_batch_id = str(uuid.uuid4())
        metadata_json = parse_json(active_version.get("metadata_json"))
        metadata_json.update(
            {
                "version": version,
                "import_batch_id": import_batch_id,
                "published_version": version,
                "is_published": False,
                "needs_republish": True,
                "last_publish_error": None,
            }
        )
        chunks = normalize_chunks_for_version(
            load_chunks(Path(active_version["chunk_json_path"])),
            metadata_json,
            {"version": version, "import_batch_id": import_batch_id},
        )
        chunk_json_path = self._write_chunks(
            document_id=document_id,
            import_batch_id=import_batch_id,
            chunks=chunks,
        )
        version_id = str(uuid.uuid4())
        connection.execute(
            """
            INSERT INTO document_versions (
                id, document_id, version, import_batch_id, raw_docx_path,
                preprocessed_text_path, chunk_json_path, metadata_json,
                status, created_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                version_id,
                document_id,
                version,
                import_batch_id,
                active_version.get("raw_docx_path"),
                active_version.get("preprocessed_text_path"),
                str(chunk_json_path),
                json.dumps(metadata_json, ensure_ascii=False),
                STATUS_READY_FOR_REVIEW,
                now,
            ),
        )
        active_relations = connection.execute(
            """
            SELECT *
            FROM document_relations
            WHERE import_batch_id = ?
            ORDER BY created_at ASC
            """,
            (active_version["import_batch_id"],),
        ).fetchall()
        for relation in active_relations:
            connection.execute(
                """
                INSERT INTO document_relations (
                    id, source_document_id, target_document_id,
                    target_document_number, relation_type, source_text,
                    import_batch_id, is_published, created_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, 0, ?)
                """,
                (
                    str(uuid.uuid4()),
                    relation["source_document_id"],
                    relation["target_document_id"],
                    relation["target_document_number"],
                    relation["relation_type"],
                    relation["source_text"],
                    import_batch_id,
                    now,
                ),
            )
        return {
            **active_version,
            "id": version_id,
            "version": version,
            "import_batch_id": import_batch_id,
            "chunk_json_path": str(chunk_json_path),
            "metadata_json": json.dumps(metadata_json, ensure_ascii=False),
            "status": STATUS_READY_FOR_REVIEW,
            "created_at": now,
        }

    def _rewrite_version_chunks_metadata(
        self,
        version: dict[str, Any],
        metadata_json: dict[str, Any],
    ) -> None:
        chunk_path = Path(version["chunk_json_path"])
        chunks = normalize_chunks_for_version(load_chunks(chunk_path), metadata_json, version)
        chunk_path.write_text(json.dumps(chunks, ensure_ascii=False, indent=2), encoding="utf-8")

    def update_document_metadata(
        self,
        document_id: str,
        payload: dict[str, Any],
    ) -> dict[str, Any]:
        allowed_fields = METADATA_UPDATE_FIELDS - {"source_url"}
        updates = {key: optional_str(payload.get(key)) for key in allowed_fields if key in payload}
        if not updates:
            raise DocumentImportError(
                "METADATA_INCOMPLETE",
                "No supported metadata fields were provided.",
            )

        now = utc_now_iso()
        with get_connection(self.db_path) as connection:
            registry, version = self._ensure_editable_latest_version(connection, document_id)
            if version is not None:
                metadata_json = parse_json(version.get("metadata_json"))
                metadata_json.update(updates)
                metadata_json["needs_republish"] = bool(registry.get("is_published"))
                blockers = get_publish_blockers(metadata_json)
                metadata_json["needs_review"] = bool(blockers)
                metadata_json["needs_review_fields"] = blockers
                connection.execute(
                    """
                    UPDATE document_versions
                    SET metadata_json = ?, status = ?
                    WHERE id = ?
                    """,
                    (
                        json.dumps(metadata_json, ensure_ascii=False),
                        STATUS_READY_FOR_REVIEW,
                        version["id"],
                    ),
                )
                self._rewrite_version_chunks_metadata(version, metadata_json)
                if not registry.get("is_published"):
                    assignments = ", ".join(f"{field} = ?" for field in updates)
                    connection.execute(
                        f"""
                        UPDATE document_registry
                        SET {assignments}, raw_metadata_json = ?, updated_at = ?
                        WHERE document_id = ?
                        """,
                        (
                            *updates.values(),
                            json.dumps(metadata_json, ensure_ascii=False),
                            now,
                            document_id,
                        ),
                    )
            connection.commit()
        detail = self.get_document_detail(document_id)
        assert detail is not None
        return detail

    def update_document_chunks(
        self,
        document_id: str,
        chunks: list[dict[str, Any]],
    ) -> dict[str, Any]:
        validate_chunks(chunks)
        now = utc_now_iso()
        with get_connection(self.db_path) as connection:
            registry, version = self._ensure_editable_latest_version(connection, document_id)
            chunk_path = Path(version["chunk_json_path"])
            chunk_path.parent.mkdir(parents=True, exist_ok=True)
            metadata_json = parse_json(version.get("metadata_json"))
            metadata_json["needs_republish"] = bool(registry.get("is_published"))
            chunks = normalize_chunks_for_version(chunks, metadata_json, version)
            chunk_path.write_text(json.dumps(chunks, ensure_ascii=False, indent=2), encoding="utf-8")
            connection.execute(
                """
                UPDATE document_versions
                SET metadata_json = ?, status = ?
                WHERE id = ?
                """,
                (
                    json.dumps(metadata_json, ensure_ascii=False),
                    STATUS_READY_FOR_REVIEW,
                    version["id"],
                ),
            )
            if not registry.get("is_published"):
                connection.execute(
                    """
                    UPDATE document_registry
                    SET raw_metadata_json = ?, updated_at = ?
                    WHERE document_id = ?
                    """,
                    (json.dumps(metadata_json, ensure_ascii=False), now, document_id),
                )
            connection.commit()
        detail = self.get_document_detail(document_id)
        assert detail is not None
        return detail

    def replace_document_relations(
        self,
        document_id: str,
        relations: list[dict[str, Any]],
    ) -> dict[str, Any]:
        for relation in relations:
            relation_type = optional_str(relation.get("relation_type"))
            target = optional_str(relation.get("target_document_number")) or optional_str(
                relation.get("target_document_id")
            )
            if not relation_type or not target:
                raise DocumentImportError(
                    "METADATA_INCOMPLETE",
                    "Each relation needs relation_type and a target document.",
                )

        now = utc_now_iso()
        with get_connection(self.db_path) as connection:
            registry, version = self._ensure_editable_latest_version(connection, document_id)
            validate_relation_targets_exist(connection, relations)
            connection.execute(
                "DELETE FROM document_relations WHERE import_batch_id = ?",
                (version["import_batch_id"],),
            )
            for relation in relations:
                connection.execute(
                    """
                    INSERT INTO document_relations (
                        id, source_document_id, target_document_id,
                        target_document_number, relation_type, source_text,
                        import_batch_id, is_published, created_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, 0, ?)
                    """,
                    (
                        str(uuid.uuid4()),
                        document_id,
                        optional_str(relation.get("target_document_id")),
                        optional_str(relation.get("target_document_number")),
                        str(relation.get("relation_type", "")).strip(),
                        optional_str(relation.get("source_text")),
                        version["import_batch_id"],
                        now,
                ),
            )
            metadata_json = parse_json(version.get("metadata_json"))
            metadata_json["relations_unavailable"] = not bool(relations)
            metadata_json["needs_republish"] = bool(registry.get("is_published"))
            connection.execute(
                """
                UPDATE document_versions
                SET metadata_json = ?, status = ?
                WHERE id = ?
                """,
                (
                    json.dumps(metadata_json, ensure_ascii=False),
                    STATUS_READY_FOR_REVIEW,
                    version["id"],
                ),
            )
            if not registry.get("is_published"):
                connection.execute(
                    """
                    UPDATE document_registry
                    SET raw_metadata_json = ?, updated_at = ?
                    WHERE document_id = ?
                    """,
                    (json.dumps(metadata_json, ensure_ascii=False), now, document_id),
                )
            connection.commit()
        detail = self.get_document_detail(document_id)
        assert detail is not None
        return detail

    def _parse_metadata(
        self,
        form_data: dict[str, Any],
        metadata_hints: dict[str, Any],
        text: str,
        paragraphs: list[str],
        import_batch_id: str,
    ) -> ImportMetadata:
        llm_metadata = infer_metadata_with_deepseek(
            config=self.config,
            text=text,
            paragraphs=paragraphs,
        )
        merged_hints = {
            **metadata_hints,
            **{key: value for key, value in llm_metadata.items() if value},
        }
        extraction_sources = {
            "regex": metadata_hints,
            "llm": llm_metadata,
        }
        document_number = first_present(
            form_data.get("document_number"),
            merged_hints.get("document_number"),
        )
        title = first_present(
            form_data.get("title"),
            merged_hints.get("title"),
        )
        issued_date = first_present(
            form_data.get("issued_date"),
            merged_hints.get("issued_date"),
        )
        effective_date = first_present(
            form_data.get("effective_date"),
            merged_hints.get("effective_date"),
        )
        issuing_body = first_present(
            form_data.get("issuing_body"),
            merged_hints.get("issuing_body"),
        )
        signer_title = first_present(
            form_data.get("signer_title"),
            merged_hints.get("signer_title"),
        )
        signer_name = first_present(
            form_data.get("signer_name"),
            merged_hints.get("signer_name"),
        )
        document_type = first_present(
            form_data.get("document_type"),
            merged_hints.get("document_type"),
        )

        needs_review_fields = []
        if not document_number:
            document_number = f"UNIDENTIFIED-{import_batch_id[:8]}"
            needs_review_fields.append("document_number")
        if not title:
            title = make_fallback_title(paragraphs, import_batch_id)
            needs_review_fields.append("title")
        for field_name, value in [
            ("document_type", document_type),
            ("issuing_body", issuing_body),
            ("issued_date", issued_date),
            ("effective_date", effective_date),
            ("validity_status", optional_str(form_data.get("validity_status"))),
            ("signer_title", signer_title),
            ("signer_name", signer_name),
        ]:
            if not value:
                needs_review_fields.append(field_name)

        llm_review_fields = llm_metadata.get("needs_review_fields")
        if isinstance(llm_review_fields, list):
            needs_review_fields.extend(str(field) for field in llm_review_fields)
        needs_review_fields.extend(low_confidence_fields(llm_metadata))
        needs_review_fields = sorted(set(needs_review_fields))

        relations_raw = str(form_data.get("relations_json", "") or "").strip()
        relations = parse_relations_json(relations_raw)
        if not relations:
            relations = normalize_inferred_relations(
                llm_metadata.get("relations") or []
            )
        return ImportMetadata(
            document_number=document_number,
            title=title,
            source_url=None,
            issuing_body=issuing_body,
            issued_date=issued_date,
            effective_date=effective_date,
            expiry_date=optional_str(form_data.get("expiry_date")),
            validity_status=optional_str(form_data.get("validity_status"))
            or VALIDITY_UNKNOWN,
            signer_title=signer_title,
            signer_name=signer_name,
            document_type=document_type,
            relations=relations,
            relations_unavailable=not bool(relations),
            needs_review=bool(needs_review_fields),
            needs_review_fields=needs_review_fields,
            extraction=extraction_sources,
        )

    def _validate_docx(self, file_storage: FileStorage | None) -> None:
        if file_storage is None or not file_storage.filename:
            raise DocumentImportError("DOCX_REQUIRED", "Uploaded .docx file is required.")

        filename = file_storage.filename.lower()
        if not filename.endswith(".docx"):
            raise DocumentImportError(
                "DOCX_REQUIRED",
                "Uploaded file must have a .docx extension.",
            )

        stream = file_storage.stream
        original_position = stream.tell()
        try:
            with zipfile.ZipFile(stream) as archive:
                names = set(archive.namelist())
                if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                    raise DocumentImportError(
                        "DOCX_REQUIRED",
                        "Uploaded file is not a valid DOCX archive.",
                    )
        except zipfile.BadZipFile as exc:
            raise DocumentImportError(
                "DOCX_REQUIRED",
                "Uploaded file is not a valid DOCX archive.",
            ) from exc
        finally:
            stream.seek(original_position)

    def _save_raw_docx(
        self,
        file_storage: FileStorage,
        import_batch_id: str,
    ) -> Path:
        raw_dir = Path("data/raw") / import_batch_id
        raw_dir.mkdir(parents=True, exist_ok=True)
        safe_filename = secure_filename(file_storage.filename or "document.docx")
        raw_docx_path = raw_dir / safe_filename
        file_storage.save(raw_docx_path)
        return raw_docx_path

    def _write_preprocessed_text(
        self,
        document_id: str,
        import_batch_id: str,
        text: str,
    ) -> Path:
        output_dir = Path("data/preprocessed") / import_batch_id
        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / f"{document_id}.txt"
        output_path.write_text(text, encoding="utf-8")
        return output_path

    def _write_chunks(
        self,
        document_id: str,
        import_batch_id: str,
        chunks: list[dict[str, Any]],
    ) -> Path:
        output_dir = Path("data/chunked") / import_batch_id
        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / f"{document_id}.json"
        output_path.write_text(
            json.dumps(chunks, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        return output_path

    def _chunk_text(
        self,
        text: str,
        document_id: str,
        import_batch_id: str,
        metadata: dict[str, Any],
        warnings: list[str],
    ) -> list[dict[str, Any]]:
        if self.config.llm_chunking_enabled:
            try:
                from src.ingestion.llm_splitter import LLMVietnameseLegalSplitter

                splitter = LLMVietnameseLegalSplitter(doc_name=document_id)
                llm_chunks = splitter.split_text(text)
                chunks = normalize_llm_chunks(
                    llm_chunks=llm_chunks,
                    document_id=document_id,
                    import_batch_id=import_batch_id,
                    metadata=metadata,
                )
                if chunks:
                    return chunks
                warnings.append("LLM chunking returned no chunks; used regex fallback.")
            except Exception as exc:
                warnings.append(f"LLM chunking failed; used regex fallback: {exc}")

        warnings.append("Used deterministic regex chunking fallback.")
        return regex_chunk_text(
            text=text,
            document_id=document_id,
            import_batch_id=import_batch_id,
            metadata=metadata,
        )

    def _upsert_document_registry(
        self,
        metadata: ImportMetadata,
        now: str,
    ) -> tuple[str, int]:
        with get_connection(self.db_path) as connection:
            row = connection.execute(
                """
                SELECT document_id
                FROM document_registry
                WHERE document_number = ? AND is_deleted = 0
                ORDER BY updated_at DESC
                LIMIT 1
                """,
                (metadata.document_number,),
            ).fetchone()

            if row is None:
                document_id = str(uuid.uuid4())
                version = 1
                connection.execute(
                    """
                    INSERT INTO document_registry (
                        document_id, document_number, title, source_system,
                        source_url, issuing_body, signer_title, signer_name,
                        document_type, issued_date, effective_date, expiry_date,
                        validity_status, raw_metadata_json, active_version,
                        is_published, is_deleted, created_at, updated_at
                    )
                    VALUES (?, ?, ?, 'admin_upload', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, NULL, 0, 0, ?, ?)
                    """,
                    (
                        document_id,
                        metadata.document_number,
                        metadata.title,
                        metadata.source_url,
                        metadata.issuing_body,
                        metadata.signer_title,
                        metadata.signer_name,
                        metadata.document_type,
                        metadata.issued_date,
                        metadata.effective_date,
                        metadata.expiry_date,
                        metadata.validity_status,
                        json.dumps(metadata.__dict__, ensure_ascii=False),
                        now,
                        now,
                    ),
                )
            else:
                document_id = row["document_id"]
                version_row = connection.execute(
                    """
                    SELECT COALESCE(MAX(version), 0) + 1 AS next_version
                    FROM document_versions
                    WHERE document_id = ?
                    """,
                    (document_id,),
                ).fetchone()
                version = int(version_row["next_version"])
                connection.execute(
                    """
                    UPDATE document_registry
                    SET title = ?,
                        source_url = ?,
                        issued_date = ?,
                        effective_date = ?,
                        expiry_date = ?,
                        validity_status = ?,
                        issuing_body = ?,
                        signer_title = ?,
                        signer_name = ?,
                        document_type = ?,
                        raw_metadata_json = ?,
                        updated_at = ?
                    WHERE document_id = ?
                    """,
                    (
                        metadata.title,
                        metadata.source_url,
                        metadata.issued_date,
                        metadata.effective_date,
                        metadata.expiry_date,
                        metadata.validity_status,
                        metadata.issuing_body,
                        metadata.signer_title,
                        metadata.signer_name,
                        metadata.document_type,
                        json.dumps(metadata.__dict__, ensure_ascii=False),
                        now,
                        document_id,
                    ),
                )
            connection.commit()

        return document_id, version

    def _store_version_and_relations(
        self,
        document_id: str,
        version: int,
        import_batch_id: str,
        raw_docx_path: Path,
        preprocessed_text_path: Path,
        chunk_json_path: Path,
        metadata_json: dict[str, Any],
        relations: list[dict[str, Any]],
        now: str,
    ) -> None:
        version_id = str(uuid.uuid4())
        with get_connection(self.db_path) as connection:
            validate_relation_targets_exist(connection, relations)
            connection.execute(
                """
                INSERT INTO document_versions (
                    id, document_id, version, import_batch_id, raw_docx_path,
                    preprocessed_text_path, chunk_json_path, metadata_json,
                    status, created_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    version_id,
                    document_id,
                    version,
                    import_batch_id,
                    str(raw_docx_path),
                    str(preprocessed_text_path),
                    str(chunk_json_path),
                    json.dumps(metadata_json, ensure_ascii=False),
                    STATUS_READY_FOR_REVIEW,
                    now,
                ),
            )
            for relation in relations:
                connection.execute(
                    """
                    INSERT INTO document_relations (
                        id, source_document_id, target_document_id,
                        target_document_number, relation_type, source_text,
                        import_batch_id, is_published, created_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, 0, ?)
                    """,
                    (
                        str(uuid.uuid4()),
                        document_id,
                        optional_str(relation.get("target_document_id")),
                        optional_str(relation.get("target_document_number")),
                        str(relation.get("relation_type", "")).strip(),
                        optional_str(relation.get("source_text")),
                        import_batch_id,
                        now,
                    ),
                )
            connection.commit()

    def _create_pipeline_run(
        self,
        pipeline_run_id: str,
        requested_by_user_id: str,
        status: str,
        input_json: dict[str, Any],
        now: str,
    ) -> None:
        with get_connection(self.db_path) as connection:
            connection.execute(
                """
                INSERT INTO pipeline_runs (
                    id, pipeline_type, requested_by_user_id, input_json,
                    status, error_message, created_at, updated_at
                )
                VALUES (?, ?, ?, ?, ?, NULL, ?, ?)
                """,
                (
                    pipeline_run_id,
                    PIPELINE_IMPORT_DOCUMENT,
                    requested_by_user_id,
                    json.dumps(input_json, ensure_ascii=False),
                    status,
                    now,
                    now,
                ),
            )
            connection.execute(
                """
                INSERT INTO pipeline_events (
                    id, pipeline_run_id, state, message, payload_json, created_at
                )
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (
                    str(uuid.uuid4()),
                    pipeline_run_id,
                    status,
                    "Import pipeline created.",
                    json.dumps(input_json, ensure_ascii=False),
                    now,
                ),
            )
            connection.commit()

    def _mark_pipeline(
        self,
        pipeline_run_id: str,
        status: str,
        message: str,
        payload: dict[str, Any] | None = None,
    ) -> None:
        now = utc_now_iso()
        with get_connection(self.db_path) as connection:
            connection.execute(
                """
                UPDATE pipeline_runs
                SET status = ?, updated_at = ?
                WHERE id = ?
                """,
                (status, now, pipeline_run_id),
            )
            connection.execute(
                """
                INSERT INTO pipeline_events (
                    id, pipeline_run_id, state, message, payload_json, created_at
                )
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (
                    str(uuid.uuid4()),
                    pipeline_run_id,
                    status,
                    message,
                    json.dumps(payload or {}, ensure_ascii=False),
                    now,
                ),
            )
            connection.commit()

    def _fail_pipeline(
        self,
        pipeline_run_id: str,
        message: str,
        payload: dict[str, Any],
    ) -> None:
        now = utc_now_iso()
        with get_connection(self.db_path) as connection:
            connection.execute(
                """
                UPDATE pipeline_runs
                SET status = ?, error_message = ?, updated_at = ?
                WHERE id = ?
                """,
                (STATUS_FAILED, message, now, pipeline_run_id),
            )
            connection.execute(
                """
                INSERT INTO pipeline_events (
                    id, pipeline_run_id, state, message, payload_json, created_at
                )
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (
                    str(uuid.uuid4()),
                    pipeline_run_id,
                    STATUS_FAILED,
                    message,
                    json.dumps(payload, ensure_ascii=False),
                    now,
                ),
            )
            connection.commit()


def optional_str(value: Any) -> str | None:
    if value is None:
        return None
    stripped = str(value).strip()
    return stripped or None


def bool_from_form(value: Any, default: bool) -> bool:
    if value is None:
        return default
    return str(value).strip().lower() in {"1", "true", "yes", "on"}


def first_present(*values: Any) -> str | None:
    for value in values:
        candidate = optional_str(value)
        if candidate:
            return candidate
    return None


def parse_json(raw_value: Any) -> dict[str, Any]:
    if not raw_value:
        return {}
    if isinstance(raw_value, dict):
        return raw_value
    try:
        parsed = json.loads(str(raw_value))
    except json.JSONDecodeError:
        return {}
    return parsed if isinstance(parsed, dict) else {}


def overlay_metadata(document: dict[str, Any], metadata: dict[str, Any]) -> None:
    for field in METADATA_UPDATE_FIELDS - {"source_url"}:
        if field in metadata:
            document[field] = metadata.get(field)


def get_publish_blockers(metadata: dict[str, Any]) -> list[str]:
    blockers = [
        field
        for field in REQUIRED_PUBLISH_FIELDS
        if not optional_str(metadata.get(field))
    ]
    if optional_str(metadata.get("validity_status")) == VALIDITY_UNKNOWN:
        blockers.append("validity_status")
    return sorted(set(blockers))


def normalize_chunks_for_version(
    chunks: list[dict[str, Any]],
    metadata: dict[str, Any],
    version: dict[str, Any],
) -> list[dict[str, Any]]:
    normalized = []
    for chunk in chunks:
        updated = dict(chunk)
        updated["import_batch_id"] = str(version["import_batch_id"])
        updated["is_published"] = False
        updated["published_version"] = int(version["version"])
        updated["document_number"] = metadata.get("document_number") or updated.get("document_number")
        updated["document_title"] = metadata.get("title") or updated.get("document_title")
        updated["document_type"] = metadata.get("document_type") or updated.get("document_type")
        updated["issuing_body"] = metadata.get("issuing_body") or updated.get("issuing_body")
        updated["issued_date"] = metadata.get("issued_date") or updated.get("issued_date")
        updated["effective_date"] = metadata.get("effective_date") or updated.get("effective_date")
        updated["expiry_date"] = metadata.get("expiry_date") or updated.get("expiry_date")
        updated["validity_status"] = metadata.get("validity_status") or updated.get("validity_status")
        normalized.append(updated)
    return normalized


def make_fallback_title(paragraphs: list[str], import_batch_id: str) -> str:
    for paragraph in paragraphs[:10]:
        if paragraph.strip():
            return paragraph.strip()[:160]
    return f"Untitled document {import_batch_id[:8]}"


def normalize_inferred_relations(raw_relations: Any) -> list[dict[str, Any]]:
    if not isinstance(raw_relations, list):
        return []
    relations = []
    for relation in raw_relations:
        if not isinstance(relation, dict):
            continue
        relation_type = optional_str(relation.get("relation_type"))
        target = optional_str(relation.get("target_document_number")) or optional_str(
            relation.get("target_document_id")
        )
        if not relation_type or not target:
            continue
        relations.append(
            {
                "relation_type": relation_type,
                "target_document_number": optional_str(
                    relation.get("target_document_number")
                ),
                "target_document_id": optional_str(relation.get("target_document_id")),
                "source_text": optional_str(relation.get("source_text")),
                "source": optional_str(relation.get("source")) or "llm",
            }
        )
    return relations


def validate_relation_targets_exist(connection, relations: list[dict[str, Any]]) -> None:
    for relation in relations:
        relation_type = optional_str(relation.get("relation_type"))
        target_document_number = optional_str(relation.get("target_document_number"))
        if not relation_type or not target_document_number:
            raise DocumentImportError(
                "RELATION_TARGET_NOT_FOUND",
                "Each relation needs relation_type and an existing target_document_number.",
                details={
                    "relation_type": relation_type,
                    "target_document_number": target_document_number,
                },
            )
        row = connection.execute(
            """
            SELECT document_id
            FROM document_registry
            WHERE document_number = ? AND is_deleted = 0
            LIMIT 1
            """,
            (target_document_number,),
        ).fetchone()
        if row is None:
            raise DocumentImportError(
                "RELATION_TARGET_NOT_FOUND",
                "Relationship target document does not exist.",
                details={
                    "relation_type": relation_type,
                    "target_document_number": target_document_number,
                },
            )


def low_confidence_fields(metadata: dict[str, Any], threshold: float = 0.7) -> list[str]:
    confidence = metadata.get("confidence")
    if not isinstance(confidence, dict):
        return []
    fields = []
    for field, score in confidence.items():
        try:
            numeric_score = float(score)
        except (TypeError, ValueError):
            fields.append(str(field))
            continue
        if numeric_score < threshold:
            fields.append(str(field))
    return fields


def parse_relations_json(raw_value: str) -> list[dict[str, Any]]:
    if not raw_value:
        return []
    try:
        parsed = json.loads(raw_value)
    except json.JSONDecodeError as exc:
        raise DocumentImportError(
            "METADATA_INCOMPLETE",
            "relations_json must be valid JSON when provided.",
            details={"relations_json_error": str(exc)},
        ) from exc

    if isinstance(parsed, dict):
        parsed = parsed.get("relations", [])
    if not isinstance(parsed, list):
        raise DocumentImportError(
            "METADATA_INCOMPLETE",
            "relations_json must be a list or an object with a relations list.",
        )

    relations = []
    for item in parsed:
        if not isinstance(item, dict):
            raise DocumentImportError(
                "METADATA_INCOMPLETE",
                "Each relation must be an object.",
            )
        relation_type = str(item.get("relation_type", "")).strip()
        target_document_number = optional_str(item.get("target_document_number"))
        target_document_id = optional_str(item.get("target_document_id"))
        if not relation_type or not (target_document_number or target_document_id):
            raise DocumentImportError(
                "METADATA_INCOMPLETE",
                "Each relation needs relation_type and a target document.",
            )
        relations.append(item)
    return relations


def infer_metadata_with_deepseek(
    config: Config,
    text: str,
    paragraphs: list[str],
) -> dict[str, Any]:
    api_key = os.getenv("DEEPSEEK_API_KEY")
    if not api_key:
        return {}
    try:
        from openai import OpenAI
    except ImportError:
        return {}

    prompt = {
        "first_context": make_first_context(paragraphs, text),
        "effect_context": extract_effective_clause_text(text),
        "closing_context": make_closing_context(paragraphs, text),
        "instructions": (
            "Extract Vietnamese legal document metadata. Return JSON only with keys: "
            "title, document_number, issuing_body, issued_date, effective_date, signer_name, "
            "signer_title, document_type, relations, confidence, needs_review_fields, evidence. "
            "confidence must be an object with field names mapped to 0.0-1.0 scores. "
            "Dates must use YYYY-MM-DD. Use null when unsure."
        ),
    }
    try:
        client = OpenAI(api_key=api_key, base_url="https://api.deepseek.com/v1")
        response = client.chat.completions.create(
            model=config.llm_model_metadata or "deepseek-v4-flash",
            messages=[
                {
                    "role": "system",
                    "content": "You extract metadata and strictly return one JSON object.",
                },
                {
                    "role": "user",
                    "content": json.dumps(prompt, ensure_ascii=False),
                },
            ],
            response_format={"type": "json_object"},
        )
        content = response.choices[0].message.content or "{}"
        parsed = json.loads(content)
        return parsed if isinstance(parsed, dict) else {}
    except Exception:
        return {}


def make_first_context(paragraphs: list[str], text: str, max_words: int = 600) -> str:
    context = "\n".join(paragraphs[:25]) or text
    words = context.split()
    return " ".join(words[:max_words])


def make_closing_context(paragraphs: list[str], text: str, max_words: int = 350) -> str:
    context = "\n".join(paragraphs[-25:]) or text[-3000:]
    words = context.split()
    return " ".join(words[-max_words:])


def extract_effective_clause_text(text: str) -> str:
    patterns = [
        r"(?is)(?:Điều|Dieu)\s+\d+[^\n\r]*hiệu lực thi hành.*?(?=(?:Điều|Dieu)\s+\d+|\Z)",
        r"(?is)(?:Điều|Dieu)\s+\d+[^\n\r]*hieu luc thi hanh.*?(?=(?:Điều|Dieu)\s+\d+|\Z)",
        r"(?is)(?:Äiá»u|Dieu)\s+\d+[^\n\r]*hi[eÃª]u l[uÆ°]c thi h[aÃ ]nh.*?(?=(?:Äiá»u|Dieu)\s+\d+|\Z)",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return match.group(0).strip()[:3000]
    return ""


def extract_effective_date(text: str) -> str | None:
    for pattern in [
        r"(\d{1,2})/(\d{1,2})/(\d{4})",
        r"(?:ngay|ngày)\s+(\d{1,2})\s+(?:thang|tháng)\s+(\d{1,2})\s+(?:nam|năm)\s+(\d{4})",
    ]:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            day, month, year = match.groups()
            return f"{int(year):04d}-{int(month):02d}-{int(day):02d}"
    return None


def extract_signature(
    paragraphs: list[str],
    tables: list[dict[str, Any]] | None = None,
) -> tuple[str | None, str | None]:
    table_signature_lines = extract_signature_lines_from_last_table(tables or [])
    tail = [
        paragraph.strip()
        for paragraph in [*paragraphs[-35:], *table_signature_lines]
        if paragraph.strip()
    ]
    signer_title = None
    signer_name = None
    title_markers = (
        "BO TRUONG",
        "THU TRUONG",
        "CHU TICH",
        "THU TUONG",
        "PHO THU TUONG",
        "KT.",
        "TM.",
        "BỘ TRƯỞNG",
        "THỦ TRƯỞNG",
        "CHỦ TỊCH",
        "THỦ TƯỚNG",
        "PHÓ THỦ TƯỚNG",
    )
    for index, paragraph in enumerate(tail):
        if looks_like_signer_title(paragraph):
            signer_title = paragraph
            following = tail[index + 1 : index + 8]
            for candidate in reversed(following):
                if looks_like_signer_name(candidate):
                    signer_name = candidate
                    break
            break
    if signer_name is None:
        for candidate in reversed(tail):
            if looks_like_signer_name(candidate):
                signer_name = candidate
                break
    return signer_title, signer_name


def extract_signature_lines_from_last_table(tables: list[dict[str, Any]]) -> list[str]:
    if not tables:
        return []
    for table in reversed(tables):
        lines = flatten_table_text(table)
        if any(looks_like_signer_title(line) for line in lines) or any(
            looks_like_signer_name(line) for line in lines
        ):
            return lines
    return []


def flatten_table_text(table: dict[str, Any]) -> list[str]:
    lines: list[str] = []
    for row in table.get("rows", []):
        for cell in row:
            for line in str(cell).splitlines():
                cleaned = line.strip()
                if cleaned:
                    lines.append(cleaned)
    return dedupe_preserving_order(lines)


def dedupe_preserving_order(values: list[str]) -> list[str]:
    seen = set()
    deduped = []
    for value in values:
        normalized = re.sub(r"\s+", " ", value).strip().upper()
        if not normalized or normalized in seen:
            continue
        seen.add(normalized)
        deduped.append(value)
    return deduped


def looks_like_signer_title(value: str) -> bool:
    normalized = value.upper()
    title_markers = (
        "BO TRUONG",
        "THU TRUONG",
        "CHU TICH",
        "CHU TICH QUOC HOI",
        "THU TUONG",
        "PHO THU TUONG",
        "KT.",
        "TM.",
        "BỘ TRƯỞNG",
        "THỦ TRƯỞNG",
        "CHỦ TỊCH",
        "CHỦ TỊCH QUỐC HỘI",
        "THỦ TƯỚNG",
        "PHÓ THỦ TƯỚNG",
        "Bá»˜ TRÆ¯á»žNG",
        "THá»¦ TRÆ¯á»žNG",
        "CHá»¦ Tá»ŠCH",
        "THá»¦ TÆ¯á»šNG",
        "PHÃ“ THá»¦ TÆ¯á»šNG",
    )
    return any(marker in normalized for marker in title_markers)


def looks_like_signer_name(value: str) -> bool:
    words = value.split()
    if not (2 <= len(words) <= 6):
        return False
    upper = value.upper()
    if any(char.isdigit() for char in value):
        return False
    if upper.startswith(("NOI NHAN", "NƠI NHẬN", "DIEU", "ĐIỀU")):
        return False
    uppercase_letters = sum(1 for char in value if char.isalpha() and char.isupper())
    letters = sum(1 for char in value if char.isalpha())
    if letters > 0 and uppercase_letters / letters >= 0.4:
        return True
    return all(word[:1].isupper() for word in words)


def extract_document_type_from_title(title: Any) -> str | None:
    if not title:
        return None
    normalized = str(title).strip().upper()
    for prefix, document_type in [
        ("BỘ LUẬT", "Bộ luật"),
        ("BO LUAT", "Bộ luật"),
        ("LUẬT", "Luật"),
        ("LUAT", "Luật"),
        ("NGHỊ ĐỊNH", "Nghị định"),
        ("NGHI DINH", "Nghị định"),
        ("THÔNG TƯ", "Thông tư"),
        ("THONG TU", "Thông tư"),
        ("QUYẾT ĐỊNH", "Quyết định"),
        ("QUYET DINH", "Quyết định"),
        ("NGHỊ QUYẾT", "Nghị quyết"),
        ("NGHI QUYET", "Nghị quyết"),
    ]:
        if normalized.startswith(prefix):
            return document_type
    return None


def extract_docx(docx_path: Path) -> dict[str, Any]:
    document = Document(docx_path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
    paragraphs = [text for text in paragraphs if text]
    tables = []
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = []
            seen = set()
            for cell in row.cells:
                text = cell.text.strip()
                if text and text not in seen:
                    cells.append(text)
                    seen.add(text)
            if cells:
                rows.append(cells)
        tables.append({"table_index": table_index, "rows": rows})
    return {"paragraphs": paragraphs, "tables": tables}


def extract_metadata_hints(
    text: str,
    tables: list[dict[str, Any]],
    paragraphs: list[str],
) -> dict[str, Any]:
    head = "\n".join(text.splitlines()[:80])
    hints: dict[str, Any] = {}

    issuing_body = extract_issuing_body_from_first_table(tables)
    if issuing_body:
        hints["issuing_body"] = issuing_body

    document_number = extract_document_number_from_first_table(tables)
    if document_number:
        hints["document_number"] = document_number
    else:
        document_number_match = re.search(
            r"(?:Số|So)\s*:\s*([^\n\r]+)",
            head,
            flags=re.IGNORECASE,
        )
        if document_number_match:
            hints["document_number"] = cleanup_document_number(
                document_number_match.group(1)
            )

    date_match = re.search(
        r"(?:ngày|ngay)\s+(\d{1,2})\s+"
        r"(?:tháng|thang)\s+(\d{1,2})\s+"
        r"(?:năm|nam)\s+(\d{4})",
        head,
        flags=re.IGNORECASE,
    )
    if date_match:
        day, month, year = date_match.groups()
        hints["issued_date"] = f"{int(year):04d}-{int(month):02d}-{int(day):02d}"

    title = extract_title_from_paragraphs(paragraphs)
    if title:
        hints["title"] = title
    else:
        for line in text.splitlines()[:30]:
            normalized = line.strip()
            if re.match(r"^(LUẬT|NGHỊ ĐỊNH|THÔNG TƯ|QUYẾT ĐỊNH)\b", normalized, re.I):
                hints["title"] = normalized
                break

    effective_text = extract_effective_clause_text(text)
    effective_date = extract_effective_date(effective_text or text)
    if effective_date:
        hints["effective_date"] = effective_date
    signer_title, signer_name = extract_signature(paragraphs, tables)
    if signer_title:
        hints["signer_title"] = signer_title
    if signer_name:
        hints["signer_name"] = signer_name
    document_type = extract_document_type_from_title(hints.get("title"))
    if document_type:
        hints["document_type"] = document_type

    return hints


def extract_issuing_body_from_first_table(
    tables: list[dict[str, Any]],
) -> str | None:
    if not tables:
        return None

    rows = tables[0].get("rows", [])
    for row_index, row in enumerate(rows[:5]):
        for cell_index, cell in enumerate(row):
            lines = split_cell_lines(str(cell))
            number_line_index = next(
                (
                    index
                    for index, line in enumerate(lines)
                    if extract_document_number_from_text(line)
                    or re.match(r"^(Số|So)\b", line.strip(), flags=re.IGNORECASE)
                ),
                None,
            )
            if number_line_index is None:
                continue
            candidates = lines[:number_line_index]
            for previous_row in reversed(rows[:row_index]):
                if cell_index < len(previous_row):
                    candidates = split_cell_lines(str(previous_row[cell_index])) + candidates
            for candidate in reversed(candidates):
                cleaned = cleanup_issuing_body(candidate)
                if is_issuing_body_candidate(cleaned):
                    return cleaned

    for row in rows[:3]:
        for cell in row[:1]:
            for line in split_cell_lines(str(cell)):
                cleaned = cleanup_issuing_body(line)
                if is_issuing_body_candidate(cleaned):
                    return cleaned
    return None


def split_cell_lines(value: str) -> list[str]:
    return [line.strip() for line in value.splitlines() if line.strip()]


def cleanup_issuing_body(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip(" .;,:")


def is_issuing_body_candidate(value: str | None) -> bool:
    if not value:
        return False
    normalized = value.upper()
    if len(value) > 160 or any(char.isdigit() for char in value):
        return False
    if extract_document_number_from_text(value):
        return False
    excluded_markers = (
        "CONG HOA",
        "CỘNG HÒA",
        "DOC LAP",
        "ĐỘC LẬP",
        "SOCIALIST",
        "VIET NAM",
        "Số",
        "SO:",
    )
    if any(marker.upper() in normalized for marker in excluded_markers):
        return False
    body_markers = (
        "QUOC HOI",
        "QUỐC HỘI",
        "CHINH PHU",
        "CHÍNH PHỦ",
        "BO ",
        "BỘ ",
        "TOA AN",
        "TÒA ÁN",
        "VIEN KIEM SAT",
        "VIỆN KIỂM SÁT",
        "UY BAN",
        "ỦY BAN",
        "HOI DONG",
        "HỘI ĐỒNG",
        "NGAN HANG",
        "NGÂN HÀNG",
        "KIEM TOAN",
        "KIỂM TOÁN",
    )
    return any(marker in normalized for marker in body_markers)


def extract_document_number_from_first_table(
    tables: list[dict[str, Any]],
) -> str | None:
    if not tables:
        return None

    first_table = tables[0]
    for row in first_table.get("rows", []):
        for cell in row:
            document_number = extract_document_number_from_text(str(cell))
            if document_number:
                return document_number

    table_text = "\n".join(
        str(cell)
        for row in first_table.get("rows", [])
        for cell in row
    )
    return extract_document_number_from_text(table_text)


def extract_document_number_from_text(value: str) -> str | None:
    patterns = [
        r"(?:Số|So)\s*[:：]\s*([^\n\r]+)",
        r"(?:Số|So)\s+([0-9][^\n\r]+)",
    ]
    for pattern in patterns:
        match = re.search(pattern, value, flags=re.IGNORECASE)
        if match:
            return cleanup_document_number(match.group(1))
    return None


def cleanup_document_number(value: str) -> str:
    cleaned = re.split(r"\s{2,}|\t|\|", value.strip())[0].strip()
    return cleaned.strip(" .;,")


def extract_title_from_paragraphs(paragraphs: list[str]) -> str | None:
    non_empty = [paragraph.strip() for paragraph in paragraphs if paragraph.strip()]
    if not non_empty:
        return None

    first = non_empty[0]
    if not is_standalone_document_type(first):
        return first

    title_parts = [first]
    for paragraph in non_empty[1:5]:
        if is_title_boundary(paragraph):
            break
        title_parts.append(paragraph)
        if len(title_parts) >= 3:
            break

    return " ".join(title_parts)


def is_standalone_document_type(value: str) -> bool:
    normalized = value.strip(" .:;").upper()
    return normalized in {
        "LUẬT",
        "BỘ LUẬT",
        "NGHỊ ĐỊNH",
        "THÔNG TƯ",
        "QUYẾT ĐỊNH",
        "NGHỊ QUYẾT",
        "LUAT",
        "BO LUAT",
        "NGHI DINH",
        "THONG TU",
        "QUYET DINH",
        "NGHI QUYET",
    }


def is_title_boundary(value: str) -> bool:
    return bool(
        re.match(
            r"^(Căn cứ|Can cu|Điều\s+\d+|Dieu\s+\d+|Chương\b|Chuong\b)",
            value.strip(),
            flags=re.IGNORECASE,
        )
    )


def build_normalized_metadata(
    document_id: str,
    version: int,
    import_batch_id: str,
    metadata: ImportMetadata,
    metadata_hints: dict[str, Any],
) -> dict[str, Any]:
    return {
        "document_id": document_id,
        "document_number": metadata.document_number,
        "document_title": metadata.title,
        "title": metadata.title,
        "source_system": "admin_upload",
        "source_url": metadata.source_url,
        "issuing_body": metadata.issuing_body,
        "issued_date": metadata.issued_date or metadata_hints.get("issued_date"),
        "effective_date": metadata.effective_date,
        "expiry_date": metadata.expiry_date,
        "validity_status": metadata.validity_status,
        "document_type": metadata.document_type,
        "signer_title": metadata.signer_title,
        "signer_name": metadata.signer_name,
        "version": version,
        "published_version": version,
        "import_batch_id": import_batch_id,
        "is_published": False,
    }


def normalize_llm_chunks(
    llm_chunks: list[dict[str, Any]],
    document_id: str,
    import_batch_id: str,
    metadata: dict[str, Any],
) -> list[dict[str, Any]]:
    normalized = []
    for index, chunk in enumerate(llm_chunks, start=1):
        content = str(chunk.get("content", "")).strip()
        if not content:
            continue
        raw_meta = chunk.get("metadata") or {}
        article_number = (
            optional_str(raw_meta.get("article_number"))
            or optional_str(raw_meta.get("article"))
            or str(index)
        )
        clause_number = optional_str(raw_meta.get("clause_number")) or optional_str(
            raw_meta.get("clause")
        )
        chunk_level = "clause" if clause_number else "article"
        normalized.append(
            build_chunk_record(
                content=content,
                document_id=document_id,
                import_batch_id=import_batch_id,
                metadata=metadata,
                article_number=article_number,
                clause_number=clause_number,
                chunk_level=chunk_level,
                ordinal=index,
            )
        )
    return normalized


def regex_chunk_text(
    text: str,
    document_id: str,
    import_batch_id: str,
    metadata: dict[str, Any],
) -> list[dict[str, Any]]:
    article_matches = list(
        re.finditer(
            r"(?mi)^\s*(?:Điều|Dieu)\s+([0-9]+[a-zA-Z]?)\.?\s*(.*)$",
            text,
        )
    )
    if not article_matches:
        return [
            build_chunk_record(
                content=text.strip(),
                document_id=document_id,
                import_batch_id=import_batch_id,
                metadata=metadata,
                article_number="1",
                clause_number=None,
                chunk_level="article",
                ordinal=1,
            )
        ]

    chunks: list[dict[str, Any]] = []
    for article_index, match in enumerate(article_matches, start=1):
        start = match.start()
        end = (
            article_matches[article_index].start()
            if article_index < len(article_matches)
            else len(text)
        )
        article_text = text[start:end].strip()
        article_number = match.group(1)
        clause_chunks = split_article_clauses(article_text)
        if not clause_chunks:
            chunks.append(
                build_chunk_record(
                    content=article_text,
                    document_id=document_id,
                    import_batch_id=import_batch_id,
                    metadata=metadata,
                    article_number=article_number,
                    clause_number=None,
                    chunk_level="article",
                    ordinal=len(chunks) + 1,
                )
            )
            continue
        for clause_number, clause_text in clause_chunks:
            chunks.append(
                build_chunk_record(
                    content=clause_text,
                    document_id=document_id,
                    import_batch_id=import_batch_id,
                    metadata=metadata,
                    article_number=article_number,
                    clause_number=clause_number,
                    chunk_level="clause",
                    ordinal=len(chunks) + 1,
                )
            )
    return chunks


def split_article_clauses(article_text: str) -> list[tuple[str, str]]:
    matches = list(re.finditer(r"(?m)^\s*(\d+)\.\s+", article_text))
    if len(matches) < 2:
        return []
    clauses = []
    for index, match in enumerate(matches):
        start = match.start()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(article_text)
        clauses.append((match.group(1), article_text[start:end].strip()))
    return clauses


def build_chunk_record(
    content: str,
    document_id: str,
    import_batch_id: str,
    metadata: dict[str, Any],
    article_number: str,
    clause_number: str | None,
    chunk_level: str,
    ordinal: int,
) -> dict[str, Any]:
    chunk_id = make_chunk_id(document_id, article_number, clause_number, ordinal)
    citation_label = make_citation_label(
        metadata.get("document_number"),
        article_number,
        clause_number,
    )
    hierarchy_path = (
        f"Document/{metadata.get('document_number')}/Article/{article_number}"
        + (f"/Clause/{clause_number}" if clause_number else "")
    )
    content = format_clause_content(
        content=content,
        article_number=article_number,
        clause_number=clause_number,
        chunk_level=chunk_level,
    )
    return {
        "chunk_id": chunk_id,
        "document_id": document_id,
        "import_batch_id": import_batch_id,
        "document_number": metadata.get("document_number"),
        "document_title": metadata.get("document_title"),
        "document_type": metadata.get("document_type"),
        "source_system": "admin_upload",
        "source_url": metadata.get("source_url"),
        "content": content,
        "article_number": article_number,
        "article_title": None,
        "clause_number": clause_number,
        "citation_label": citation_label,
        "hierarchy_path": hierarchy_path,
        "chunk_level": chunk_level,
        "validity_status": metadata.get("validity_status", VALIDITY_UNKNOWN),
        "effective_date": metadata.get("effective_date"),
        "expiry_date": metadata.get("expiry_date"),
        "is_published": False,
        "published_version": metadata.get("published_version"),
        "ordinal": ordinal,
    }


def format_clause_content(
    content: str,
    article_number: str,
    clause_number: str | None,
    chunk_level: str,
) -> str:
    content = content.strip()
    if chunk_level != "clause" or not clause_number:
        return content

    prefix = f"Điều {article_number}.{clause_number}. "
    if content.startswith(prefix.rstrip()):
        return content
    already_contextualized_pattern = (
        rf"^\s*Điều\s+{re.escape(article_number)}\s*[,\.]\s*"
        rf"(?:khoản\s+)?{re.escape(clause_number)}\b"
    )
    if re.match(already_contextualized_pattern, content, re.IGNORECASE):
        return content
    if re.match(rf"^\s*{re.escape(clause_number)}\.\s+", content):
        return prefix + content.split(".", 1)[1].lstrip()
    return prefix + content


def make_chunk_id(
    document_id: str,
    article_number: str,
    clause_number: str | None,
    ordinal: int,
) -> str:
    raw = f"{document_id}:{article_number}:{clause_number or ''}:{ordinal}"
    return str(uuid.uuid5(uuid.NAMESPACE_URL, raw))


def make_citation_label(
    document_number: str | None,
    article_number: str,
    clause_number: str | None,
) -> str:
    label = f"{document_number or 'unknown'}, Điều {article_number}"
    if clause_number:
        label += f", khoản {clause_number}"
    return label


def validate_chunks(chunks: list[dict[str, Any]]) -> None:
    if not chunks:
        raise DocumentImportError(
            "CHUNK_VALIDATION_FAILED",
            "No valid chunks were produced from DOCX text.",
        )
    for index, chunk in enumerate(chunks, start=1):
        if not str(chunk.get("content", "")).strip():
            raise DocumentImportError(
                "CHUNK_VALIDATION_FAILED",
                f"Chunk #{index} has empty content.",
            )
        if not chunk.get("article_number"):
            raise DocumentImportError(
                "CHUNK_VALIDATION_FAILED",
                f"Chunk #{index} is missing article_number.",
            )
        if not chunk.get("hierarchy_path"):
            raise DocumentImportError(
                "CHUNK_VALIDATION_FAILED",
                f"Chunk #{index} is missing hierarchy_path.",
            )
        if chunk.get("chunk_level") == "clause" and not chunk.get("clause_number"):
            raise DocumentImportError(
                "CHUNK_VALIDATION_FAILED",
                f"Clause chunk #{index} is missing clause_number.",
            )


def load_chunk_preview(path: Path, limit: int) -> list[dict[str, Any]]:
    chunks = load_chunks(path)
    preview = []
    for chunk in chunks[:limit]:
        preview.append(
            {
                "chunk_id": chunk.get("chunk_id"),
                "citation_label": chunk.get("citation_label"),
                "chunk_level": chunk.get("chunk_level"),
                "content": str(chunk.get("content", ""))[:1000],
            }
        )
    return preview


def load_chunks(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    try:
        chunks = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return []
    return chunks if isinstance(chunks, list) else []


def count_chunks(path: Path) -> int:
    return len(load_chunks(path))


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()
