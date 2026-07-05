from __future__ import annotations

import io
import json
import os
import sqlite3
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from backend.config import Config
from backend.models.database import init_db
from backend.services.auth_service import AuthService, ROLE_BUSINESS_USER
from scripts.reset_documents_and_indexes import reset_sqlite_documents


class AdminDocxImportTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.old_cwd = os.getcwd()
        os.chdir(self.temp_dir.name)
        self.db_path = str(Path(self.temp_dir.name) / "app.sqlite3")
        self.config = Config(
            app_env="test",
            sqlite_db_path=self.db_path,
            admin_username="admin",
            admin_password="password",
            flask_secret_key="test-secret",
            llm_chunking_enabled=False,
        )
        from backend.app import create_app

        self.app = create_app(self.config)
        self.client = self.app.test_client()

    def tearDown(self) -> None:
        os.chdir(self.old_cwd)
        self.temp_dir.cleanup()

    def test_admin_import_docx_reaches_ready_for_review(self) -> None:
        token = self._login("admin", "password")
        response = self.client.post(
            "/api/v1/admin/documents/import",
            headers=self._auth_headers(token),
            data={
                "file": (make_docx_file(), "sample.docx"),
                "document_number": "01/2026/QD-TEST",
                "title": "Quyet dinh test",
                "validity_status": "unknown",
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 201, response.get_data(as_text=True))
        payload = response.get_json()
        self.assertEqual(payload["status"], "ready_for_review")
        self.assertTrue(Path(payload["raw_docx_path"]).exists())
        self.assertTrue(Path(payload["preprocessed_text_path"]).exists())
        self.assertTrue(Path(payload["chunk_json_path"]).exists())
        self.assertTrue(payload["relations_unavailable"])

        detail = self.client.get(
            f"/api/v1/admin/documents/{payload['document_id']}",
            headers=self._auth_headers(token),
        )
        self.assertEqual(detail.status_code, 200, detail.get_data(as_text=True))
        detail_payload = detail.get_json()
        self.assertGreaterEqual(len(detail_payload["chunks_preview"]), 1)
        self.assertGreaterEqual(len(detail_payload["pipeline_events"]), 4)

    def test_import_extracts_document_number_and_title_from_docx(self) -> None:
        token = self._login("admin", "password")
        response = self.client.post(
            "/api/v1/admin/documents/import",
            headers=self._auth_headers(token),
            data={
                "file": (make_docx_file(), "sample.docx"),
                "validity_status": "unknown",
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 201, response.get_data(as_text=True))
        payload = response.get_json()
        self.assertEqual(payload["document_number"], "01/2026/QD-TEST")
        self.assertEqual(payload["title"], "QUYET DINH TEST")

        detail = self.client.get(
            f"/api/v1/admin/documents/{payload['document_id']}",
            headers=self._auth_headers(token),
        )
        self.assertEqual(detail.status_code, 200, detail.get_data(as_text=True))
        document = detail.get_json()["document"]
        self.assertEqual(document["document_number"], "01/2026/QD-TEST")
        self.assertEqual(document["title"], "QUYET DINH TEST")

    def test_clause_chunks_include_article_and_clause_prefix(self) -> None:
        from backend.services.document_import_service import (
            build_chunk_record,
            regex_chunk_text,
        )

        metadata = {
            "document_number": "01/2026/QD-TEST",
            "document_title": "Quyet dinh test",
            "validity_status": "active",
        }
        chunks = regex_chunk_text(
            "Dieu 5. Trach nhiem\n1. Nguoi X co nghia vu Y.\n2. Nguoi Z co quyen W.",
            document_id="doc-1",
            import_batch_id="batch-1",
            metadata=metadata,
        )

        self.assertEqual(chunks[0]["content"], "Điều 5.1. Nguoi X co nghia vu Y.")
        self.assertEqual(chunks[1]["content"], "Điều 5.2. Nguoi Z co quyen W.")

        already_prefixed = build_chunk_record(
            content="Điều 5.1. Nguoi X co nghia vu Y.",
            document_id="doc-1",
            import_batch_id="batch-1",
            metadata=metadata,
            article_number="5",
            clause_number="1",
            chunk_level="clause",
            ordinal=1,
        )
        self.assertEqual(already_prefixed["content"], "Điều 5.1. Nguoi X co nghia vu Y.")

        article = build_chunk_record(
            content="Dieu 6. Noi dung khong co khoan.",
            document_id="doc-1",
            import_batch_id="batch-1",
            metadata=metadata,
            article_number="6",
            clause_number=None,
            chunk_level="article",
            ordinal=1,
        )
        self.assertEqual(article["content"], "Dieu 6. Noi dung khong co khoan.")

    def test_import_extracts_signature_from_last_table(self) -> None:
        token = self._login("admin", "password")
        response = self.client.post(
            "/api/v1/admin/documents/import",
            headers=self._auth_headers(token),
            data={
                "file": (make_docx_file_with_signature_table(), "signature.docx"),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 201, response.get_data(as_text=True))
        document_id = response.get_json()["document_id"]
        detail = self.client.get(
            f"/api/v1/admin/documents/{document_id}",
            headers=self._auth_headers(token),
        )
        self.assertEqual(detail.status_code, 200, detail.get_data(as_text=True))
        document = detail.get_json()["document"]
        self.assertEqual(document["issuing_body"], "QUOC HOI")
        self.assertEqual(document["signer_title"], "CHU TICH QUOC HOI")
        self.assertEqual(document["signer_name"], "Tran Thanh Man")

    def test_import_with_required_metadata_auto_publishes(self) -> None:
        token = self._login("admin", "password")
        writers = make_fake_writers()
        data = {
            "file": (make_docx_file(), "sample.docx"),
            **complete_publish_metadata(),
        }
        with patch(
            "backend.services.indexing_service.build_default_indexing_providers",
            return_value=writers,
        ):
            response = self.client.post(
                "/api/v1/admin/documents/import",
                headers=self._auth_headers(token),
                data=data,
                content_type="multipart/form-data",
            )
            self.assertEqual(response.status_code, 201, response.get_data(as_text=True))
            payload = response.get_json()
            self.assertEqual(payload["status"], "queued")
            publish_result = payload["publish_result"]
            self.assertEqual(publish_result["status"], "queued")
            self.assertFalse(payload["publish_blockers"])
            self._wait_for_pipeline(token, publish_result["pipeline_run_id"], "published")

        for writer in writers:
            self.assertEqual(writer.indexed_batches, [payload["import_batch_id"]])

        detail = self.client.get(
            f"/api/v1/admin/documents/{payload['document_id']}?scope=active",
            headers=self._auth_headers(token),
        )
        self.assertEqual(detail.get_json()["version"]["status"], "published")

    def test_non_admin_cannot_import(self) -> None:
        AuthService(self.db_path).create_user("user", "password", ROLE_BUSINESS_USER)
        token = self._login("user", "password")
        response = self.client.post(
            "/api/v1/admin/documents/import",
            headers=self._auth_headers(token),
            data={
                "file": (make_docx_file(), "sample.docx"),
                "document_number": "01/2026/QD-TEST",
                "title": "Quyet dinh test",
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 403)

    def test_admin_can_publish_ready_for_review_document_with_index_writers(self) -> None:
        token = self._login("admin", "password")
        import_payload = self._import_publishable_review_document(token)
        writers = make_fake_writers()

        with patch(
            "backend.services.indexing_service.build_default_indexing_providers",
            return_value=writers,
        ):
            response = self.client.post(
                f"/api/v1/admin/documents/{import_payload['document_id']}/publish",
                headers=self._auth_headers(token),
            )
            self.assertEqual(response.status_code, 202, response.get_data(as_text=True))
            payload = response.get_json()
            self.assertEqual(payload["status"], "queued")
            pipeline_detail = self._wait_for_pipeline(
                token,
                payload["pipeline_run_id"],
                "published",
            )

        published_event = pipeline_detail["events"][-1]
        published_payload = json.loads(published_event["payload_json"])
        self.assertEqual(published_payload["writer_statuses"]["neo4j"], "indexed")
        self.assertEqual(published_payload["writer_statuses"]["chroma"], "indexed")
        self.assertEqual(published_payload["writer_statuses"]["elasticsearch"], "indexed")
        self.assertIn("No admin-curated document relations", published_payload["warnings"][0])
        self.assertIn(
            "vector_indexing",
            {event["state"] for event in pipeline_detail["events"]},
        )
        for writer in writers:
            self.assertEqual(writer.indexed_batches, [import_payload["import_batch_id"]])

        detail = self.client.get(
            f"/api/v1/admin/documents/{import_payload['document_id']}",
            headers=self._auth_headers(token),
        )
        self.assertEqual(detail.status_code, 200, detail.get_data(as_text=True))
        detail_payload = detail.get_json()
        self.assertEqual(detail_payload["version"]["status"], "published")
        self.assertEqual(detail_payload["document"]["active_version"], 1)
        self.assertEqual(detail_payload["document"]["is_published"], 1)

    def test_non_admin_cannot_publish(self) -> None:
        admin_token = self._login("admin", "password")
        import_payload = self._import_sample_document(admin_token)
        AuthService(self.db_path).create_user("user", "password", ROLE_BUSINESS_USER)
        user_token = self._login("user", "password")

        response = self.client.post(
            f"/api/v1/admin/documents/{import_payload['document_id']}/publish",
            headers=self._auth_headers(user_token),
        )

        self.assertEqual(response.status_code, 403)

    def test_duplicate_publish_returns_existing_active_pipeline_run(self) -> None:
        token = self._login("admin", "password")
        import_payload = self._import_publishable_review_document(token)

        with patch("backend.services.indexing_service.submit_publish_job"):
            first = self.client.post(
                f"/api/v1/admin/documents/{import_payload['document_id']}/publish",
                headers=self._auth_headers(token),
            )
            second = self.client.post(
                f"/api/v1/admin/documents/{import_payload['document_id']}/publish",
                headers=self._auth_headers(token),
            )

        self.assertEqual(first.status_code, 202, first.get_data(as_text=True))
        self.assertEqual(second.status_code, 202, second.get_data(as_text=True))
        first_payload = first.get_json()
        second_payload = second.get_json()
        self.assertEqual(first_payload["pipeline_run_id"], second_payload["pipeline_run_id"])
        self.assertFalse(first_payload["already_running"])
        self.assertTrue(second_payload["already_running"])

    def test_publish_failure_keeps_document_ready_for_review_and_cleans_up(self) -> None:
        token = self._login("admin", "password")
        import_payload = self._import_publishable_review_document(token)
        writers = [
            FakeIndexWriter("neo4j"),
            FakeIndexWriter("chroma", fail_index=True),
            FakeIndexWriter("elasticsearch"),
        ]

        with patch(
            "backend.services.indexing_service.build_default_indexing_providers",
            return_value=writers,
        ):
            response = self.client.post(
                f"/api/v1/admin/documents/{import_payload['document_id']}/publish",
                headers=self._auth_headers(token),
            )
            self.assertEqual(response.status_code, 202, response.get_data(as_text=True))
            pipeline_detail = self._wait_for_pipeline(
                token,
                response.get_json()["pipeline_run_id"],
                "failed",
            )

        failed_payload = json.loads(pipeline_detail["events"][-1]["payload_json"])
        self.assertEqual(failed_payload["code"], "RuntimeError")
        self.assertEqual(writers[0].deleted_batches, [import_payload["import_batch_id"]])
        self.assertEqual(writers[1].deleted_batches, [])

        detail = self.client.get(
            f"/api/v1/admin/documents/{import_payload['document_id']}",
            headers=self._auth_headers(token),
        )
        self.assertEqual(detail.get_json()["version"]["status"], "ready_for_review")

    def test_unknown_validity_status_cannot_publish(self) -> None:
        token = self._login("admin", "password")
        import_payload = self._import_sample_document(token)
        response = self.client.post(
            f"/api/v1/admin/documents/{import_payload['document_id']}/publish",
            headers=self._auth_headers(token),
        )

        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json()["error"]["code"], "DOCUMENT_NOT_PUBLISHABLE")

    def test_editing_published_document_reindexes_new_version(self) -> None:
        token = self._login("admin", "password")
        import_payload = self._import_publishable_review_document(token)
        with patch(
            "backend.services.indexing_service.build_default_indexing_providers",
            return_value=make_fake_writers(),
        ):
            publish = self.client.post(
                f"/api/v1/admin/documents/{import_payload['document_id']}/publish",
                headers=self._auth_headers(token),
            )
            self.assertEqual(publish.status_code, 202, publish.get_data(as_text=True))
            self._wait_for_pipeline(token, publish.get_json()["pipeline_run_id"], "published")
        old_batch = import_payload["import_batch_id"]

        edit_writers = make_fake_writers()
        with patch(
            "backend.services.indexing_service.build_default_indexing_providers",
            return_value=edit_writers,
        ):
            edited = self.client.patch(
                f"/api/v1/admin/documents/{import_payload['document_id']}/metadata?auto_publish=1",
                headers=self._auth_headers(token),
                json={**complete_publish_metadata(), "title": "Quyet dinh da publish lai"},
            )
            self.assertEqual(edited.status_code, 200, edited.get_data(as_text=True))
            payload = edited.get_json()
            self.assertEqual(payload["version"]["status"], "ready_for_review")
            self.assertEqual(payload["publish_result"]["status"], "queued")
            self._wait_for_pipeline(token, payload["publish_result"]["pipeline_run_id"], "published")

        active = self.client.get(
            f"/api/v1/admin/documents/{import_payload['document_id']}?scope=active",
            headers=self._auth_headers(token),
        )
        active_payload = active.get_json()
        self.assertEqual(active_payload["version"]["status"], "published")
        self.assertEqual(active_payload["document"]["active_version"], 2)
        self.assertEqual(active_payload["document"]["title"], "Quyet dinh da publish lai")
        for writer in edit_writers:
            self.assertEqual(writer.deleted_batches, [old_batch])

    def test_editing_published_document_reindex_failure_keeps_old_active(self) -> None:
        token = self._login("admin", "password")
        import_payload = self._import_publishable_review_document(token)
        with patch(
            "backend.services.indexing_service.build_default_indexing_providers",
            return_value=make_fake_writers(),
        ):
            publish = self.client.post(
                f"/api/v1/admin/documents/{import_payload['document_id']}/publish",
                headers=self._auth_headers(token),
            )
            self.assertEqual(publish.status_code, 202, publish.get_data(as_text=True))
            self._wait_for_pipeline(token, publish.get_json()["pipeline_run_id"], "published")

        failing_writers = [
            FakeIndexWriter("neo4j"),
            FakeIndexWriter("chroma", fail_index=True),
            FakeIndexWriter("elasticsearch"),
        ]
        with patch(
            "backend.services.indexing_service.build_default_indexing_providers",
            return_value=failing_writers,
        ):
            edited = self.client.patch(
                f"/api/v1/admin/documents/{import_payload['document_id']}/metadata?auto_publish=1",
                headers=self._auth_headers(token),
                json={**complete_publish_metadata(), "title": "Ban sua loi index"},
            )
            self.assertEqual(edited.status_code, 200, edited.get_data(as_text=True))
            edited_payload = edited.get_json()
            self.assertEqual(edited_payload["version"]["status"], "ready_for_review")
            self.assertEqual(edited_payload["publish_result"]["status"], "queued")
            self._wait_for_pipeline(
                token,
                edited_payload["publish_result"]["pipeline_run_id"],
                "failed",
            )

        latest = self.client.get(
            f"/api/v1/admin/documents/{import_payload['document_id']}",
            headers=self._auth_headers(token),
        )
        self.assertTrue(latest.get_json()["needs_republish"])

        active = self.client.get(
            f"/api/v1/admin/documents/{import_payload['document_id']}?scope=active",
            headers=self._auth_headers(token),
        )
        active_payload = active.get_json()
        self.assertEqual(active_payload["version"]["status"], "published")
        self.assertEqual(active_payload["document"]["active_version"], 1)
        self.assertNotEqual(active_payload["document"]["title"], "Ban sua loi index")

    def test_rollback_deletes_indexed_batch_and_marks_version_rolled_back(self) -> None:
        token = self._login("admin", "password")
        import_payload = self._import_publishable_review_document(token)
        with patch(
            "backend.services.indexing_service.build_default_indexing_providers",
            return_value=make_fake_writers(),
        ):
            publish = self.client.post(
                f"/api/v1/admin/documents/{import_payload['document_id']}/publish",
                headers=self._auth_headers(token),
            )
            self.assertEqual(publish.status_code, 202, publish.get_data(as_text=True))
            self._wait_for_pipeline(token, publish.get_json()["pipeline_run_id"], "published")
        publish_payload = publish.get_json()

        rollback_writers = make_fake_writers()
        with patch(
            "backend.services.indexing_service.build_default_indexing_providers",
            return_value=rollback_writers,
        ):
            rollback = self.client.post(
                f"/api/v1/admin/pipeline/{publish_payload['pipeline_run_id']}/rollback",
                headers=self._auth_headers(token),
            )

        self.assertEqual(rollback.status_code, 200, rollback.get_data(as_text=True))
        rollback_payload = rollback.get_json()
        self.assertEqual(rollback_payload["status"], "rolled_back")
        self.assertIsNone(rollback_payload["restored_import_batch_id"])
        for writer in rollback_writers:
            self.assertEqual(writer.deleted_batches, [import_payload["import_batch_id"]])

        detail = self.client.get(
            f"/api/v1/admin/documents/{import_payload['document_id']}",
            headers=self._auth_headers(token),
        )
        detail_payload = detail.get_json()
        self.assertEqual(detail_payload["version"]["status"], "rolled_back")
        self.assertIsNone(detail_payload["document"]["active_version"])
        self.assertEqual(detail_payload["document"]["is_published"], 0)

    def test_neo4j_database_config_is_available_for_cloud_sessions(self) -> None:
        from backend.services.indexing_service import Neo4jGraphWriter

        config = Config(
            neo4j_uri="neo4j+s://example.databases.neo4j.io",
            neo4j_user="example",
            neo4j_password="secret",
            neo4j_database="example",
        )
        writer = Neo4jGraphWriter(config)

        self.assertEqual(writer._session_kwargs(), {"database": "example"})

    def test_import_missing_metadata_is_marked_for_review_and_rejects_invalid_docx(self) -> None:
        token = self._login("admin", "password")
        missing_metadata = self.client.post(
            "/api/v1/admin/documents/import",
            headers=self._auth_headers(token),
            data={"file": (make_docx_file_without_metadata(), "sample.docx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(missing_metadata.status_code, 201, missing_metadata.get_data(as_text=True))
        missing_payload = missing_metadata.get_json()
        self.assertTrue(missing_payload["needs_review"])
        self.assertTrue(missing_payload["document_number"].startswith("UNIDENTIFIED-"))
        self.assertIn("document_number", missing_payload["needs_review_fields"])
        self.assertIn("effective_date", missing_payload["needs_review_fields"])

        invalid_docx = self.client.post(
            "/api/v1/admin/documents/import",
            headers=self._auth_headers(token),
            data={
                "file": (io.BytesIO(b"not a docx"), "sample.docx"),
                "document_number": "01/2026/QD-TEST",
                "title": "Quyet dinh test",
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(invalid_docx.status_code, 400)
        self.assertEqual(invalid_docx.get_json()["error"]["code"], "DOCX_REQUIRED")

    def test_admin_can_download_and_edit_document_metadata_chunks_and_relations(self) -> None:
        token = self._login("admin", "password")
        import_payload = self._import_sample_document(token)
        document_id = import_payload["document_id"]
        self._import_document_with_number(token, "01/2025/QD-TEST")

        download = self.client.get(
            f"/api/v1/admin/documents/{document_id}/download",
            headers=self._auth_headers(token),
            buffered=True,
        )
        self.assertEqual(download.status_code, 200)
        self.assertGreater(len(download.data), 100)
        download.close()

        metadata_update = self.client.patch(
            f"/api/v1/admin/documents/{document_id}/metadata",
            headers=self._auth_headers(token),
            json={
                "title": "Quyet dinh da sua",
                "document_number": "02/2026/QD-TEST",
                "validity_status": "active",
                "effective_date": "2026-01-01",
                "signer_title": "Bo truong",
                "signer_name": "Nguyen Van A",
            },
        )
        self.assertEqual(metadata_update.status_code, 200, metadata_update.get_data(as_text=True))
        self.assertEqual(
            metadata_update.get_json()["document"]["title"],
            "Quyet dinh da sua",
        )

        detail = metadata_update.get_json()
        chunks = detail["chunks"]
        chunks[0]["content"] = "Noi dung chunk da sua."
        chunks_update = self.client.patch(
            f"/api/v1/admin/documents/{document_id}/chunks",
            headers=self._auth_headers(token),
            json={"chunks": chunks},
        )
        self.assertEqual(chunks_update.status_code, 200, chunks_update.get_data(as_text=True))
        self.assertEqual(chunks_update.get_json()["chunks"][0]["content"], "Noi dung chunk da sua.")

        relations_update = self.client.put(
            f"/api/v1/admin/documents/{document_id}/relationships",
            headers=self._auth_headers(token),
            json={
                "relations": [
                    {
                        "relation_type": "amends",
                        "target_document_number": "01/2025/QD-TEST",
                        "source_text": "Sua doi van ban cu.",
                    }
                ]
            },
        )
        self.assertEqual(relations_update.status_code, 200, relations_update.get_data(as_text=True))
        self.assertEqual(len(relations_update.get_json()["relations"]), 1)

    def test_relationship_target_document_number_must_exist(self) -> None:
        token = self._login("admin", "password")
        import_payload = self._import_sample_document(token)
        document_id = import_payload["document_id"]
        self._import_document_with_number(token, "01/2025/QD-TEST")

        created = self.client.put(
            f"/api/v1/admin/documents/{document_id}/relationships",
            headers=self._auth_headers(token),
            json={
                "relations": [
                    {
                        "relation_type": "amends",
                        "target_document_number": "01/2025/QD-TEST",
                        "source_text": "Sua doi van ban cu.",
                    }
                ]
            },
        )
        self.assertEqual(created.status_code, 200, created.get_data(as_text=True))
        self.assertEqual(len(created.get_json()["relations"]), 1)

        rejected = self.client.put(
            f"/api/v1/admin/documents/{document_id}/relationships",
            headers=self._auth_headers(token),
            json={
                "relations": [
                    {
                        "relation_type": "replaces",
                        "target_document_number": "MISSING-DOC",
                    }
                ]
            },
        )
        self.assertEqual(rejected.status_code, 400, rejected.get_data(as_text=True))
        self.assertEqual(
            rejected.get_json()["error"]["code"],
            "RELATION_TARGET_NOT_FOUND",
        )
        self.assertEqual(
            rejected.get_json()["error"]["details"]["target_document_number"],
            "MISSING-DOC",
        )

        detail = self.client.get(
            f"/api/v1/admin/documents/{document_id}",
            headers=self._auth_headers(token),
        )
        self.assertEqual(detail.status_code, 200, detail.get_data(as_text=True))
        relations = detail.get_json()["relations"]
        self.assertEqual(len(relations), 1)
        self.assertEqual(relations[0]["target_document_number"], "01/2025/QD-TEST")

    def test_import_rejects_relations_json_when_target_document_is_missing(self) -> None:
        token = self._login("admin", "password")
        response = self.client.post(
            "/api/v1/admin/documents/import",
            headers=self._auth_headers(token),
            data={
                "file": (make_docx_file(), "source.docx"),
                "document_number": "02/2026/QD-TEST",
                "title": "Quyet dinh co relation sai",
                "validity_status": "unknown",
                "relations_json": json.dumps(
                    [
                        {
                            "relation_type": "references",
                            "target_document_number": "MISSING-DOC",
                        }
                    ]
                ),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 400, response.get_data(as_text=True))
        payload = response.get_json()
        self.assertEqual(payload["error"]["code"], "RELATION_TARGET_NOT_FOUND")
        self.assertEqual(
            payload["error"]["details"]["target_document_number"],
            "MISSING-DOC",
        )

    def test_admin_can_hard_delete_document_and_keep_pipeline_logs(self) -> None:
        token = self._login("admin", "password")
        import_payload = self._import_publishable_review_document(token)
        document_id = import_payload["document_id"]
        writers = make_fake_writers()

        with patch(
            "backend.services.indexing_service.build_default_indexing_providers",
            return_value=writers,
        ):
            publish = self.client.post(
                f"/api/v1/admin/documents/{document_id}/publish",
                headers=self._auth_headers(token),
            )
            self.assertEqual(publish.status_code, 202, publish.get_data(as_text=True))
            self._wait_for_pipeline(token, publish.get_json()["pipeline_run_id"], "published")

        raw_path = Path(import_payload["raw_docx_path"])
        chunk_path = Path(import_payload["chunk_json_path"])
        delete_writers = make_fake_writers()
        with patch(
            "backend.services.indexing_service.build_default_indexing_providers",
            return_value=delete_writers,
        ):
            delete = self.client.delete(
                f"/api/v1/admin/documents/{document_id}",
                headers=self._auth_headers(token),
            )
        self.assertEqual(delete.status_code, 200, delete.get_data(as_text=True))
        self.assertEqual(delete.get_json()["status"], "deleted")
        for writer in delete_writers:
            self.assertEqual(writer.deleted_batches, [import_payload["import_batch_id"]])
        self.assertFalse(raw_path.exists())
        self.assertFalse(chunk_path.exists())

        detail = self.client.get(
            f"/api/v1/admin/documents/{document_id}",
            headers=self._auth_headers(token),
        )
        self.assertEqual(detail.status_code, 404)

        connection = sqlite3.connect(self.db_path)
        try:
            registry_count = connection.execute(
                "SELECT COUNT(*) FROM document_registry WHERE document_id = ?",
                (document_id,),
            ).fetchone()[0]
            version_count = connection.execute(
                "SELECT COUNT(*) FROM document_versions WHERE document_id = ?",
                (document_id,),
            ).fetchone()[0]
            event_count = connection.execute("SELECT COUNT(*) FROM pipeline_events").fetchone()[0]
        finally:
            connection.close()
        self.assertEqual(registry_count, 0)
        self.assertEqual(version_count, 0)
        self.assertGreater(event_count, 0)

    def test_reset_documents_script_clears_documents_and_preserves_users(self) -> None:
        token = self._login("admin", "password")
        self._import_sample_document(token)

        reset_sqlite_documents(Path(self.db_path), dry_run=False)

        connection = sqlite3.connect(self.db_path)
        try:
            users_count = connection.execute("SELECT COUNT(*) FROM users").fetchone()[0]
            documents_count = connection.execute("SELECT COUNT(*) FROM document_registry").fetchone()[0]
            versions_count = connection.execute("SELECT COUNT(*) FROM document_versions").fetchone()[0]
            document_runs_count = connection.execute(
                "SELECT COUNT(*) FROM pipeline_runs WHERE pipeline_type = 'import_document'"
            ).fetchone()[0]
        finally:
            connection.close()
        self.assertGreaterEqual(users_count, 1)
        self.assertEqual(documents_count, 0)
        self.assertEqual(versions_count, 0)
        self.assertEqual(document_runs_count, 0)

    def test_old_crawl_batch_id_is_migrated_to_import_batch_id(self) -> None:
        old_db_path = str(Path(self.temp_dir.name) / "old.sqlite3")
        connection = sqlite3.connect(old_db_path)
        try:
            connection.executescript(
                """
                CREATE TABLE document_registry (
                    document_id TEXT PRIMARY KEY,
                    document_number TEXT,
                    title TEXT,
                    source_system TEXT,
                    source_url TEXT,
                    sector TEXT,
                    domain TEXT,
                    issuing_body TEXT,
                    signer_title TEXT,
                    signer_name TEXT,
                    document_type TEXT,
                    issued_date TEXT,
                    effective_date TEXT,
                    expiry_date TEXT,
                    validity_status TEXT,
                    raw_metadata_json TEXT,
                    active_version INTEGER,
                    is_published INTEGER NOT NULL DEFAULT 0,
                    is_deleted INTEGER NOT NULL DEFAULT 0,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );
                CREATE TABLE document_versions (
                    id TEXT PRIMARY KEY,
                    document_id TEXT NOT NULL,
                    version INTEGER NOT NULL,
                    crawl_batch_id TEXT,
                    raw_docx_path TEXT,
                    preprocessed_text_path TEXT,
                    chunk_json_path TEXT,
                    metadata_json TEXT,
                    status TEXT NOT NULL,
                    created_at TEXT NOT NULL
                );
                INSERT INTO document_versions (
                    id, document_id, version, crawl_batch_id, status, created_at
                )
                VALUES ('v1', 'd1', 1, 'old-batch', 'ready_for_review', 'now');
                """
            )
        finally:
            connection.close()

        init_db(old_db_path)

        connection = sqlite3.connect(old_db_path)
        try:
            row = connection.execute(
                "SELECT import_batch_id FROM document_versions WHERE id = 'v1'"
            ).fetchone()
        finally:
            connection.close()
        self.assertEqual(row[0], "old-batch")

    def test_title_hint_combines_document_type_and_next_title_line(self) -> None:
        from backend.services.document_import_service import extract_title_from_paragraphs

        title = extract_title_from_paragraphs(
            [
                "LUẬT",
                "An ninh Quốc gia",
                "Căn cứ vào Hiến pháp nước Cộng hòa xã hội chủ nghĩa Việt Nam",
            ]
        )
        self.assertEqual(title, "LUẬT An ninh Quốc gia")

    def _login(self, username: str, password: str) -> str:
        response = self.client.post(
            "/api/v1/auth/login",
            json={"username": username, "password": password},
        )
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        return response.get_json()["access_token"]

    def _import_sample_document(self, token: str) -> dict:
        response = self.client.post(
            "/api/v1/admin/documents/import",
            headers=self._auth_headers(token),
            data={
                "file": (make_docx_file(), "sample.docx"),
                "validity_status": "unknown",
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 201, response.get_data(as_text=True))
        return response.get_json()

    def _import_document_with_number(self, token: str, document_number: str) -> dict:
        response = self.client.post(
            "/api/v1/admin/documents/import",
            headers=self._auth_headers(token),
            data={
                "file": (make_docx_file(), f"{document_number}.docx"),
                "document_number": document_number,
                "title": f"Target {document_number}",
                "validity_status": "unknown",
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 201, response.get_data(as_text=True))
        return response.get_json()

    def _import_publishable_review_document(self, token: str) -> dict:
        payload = self._import_sample_document(token)
        response = self.client.patch(
            f"/api/v1/admin/documents/{payload['document_id']}/metadata",
            headers=self._auth_headers(token),
            json=complete_publish_metadata(),
        )
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        return {
            **payload,
            "import_batch_id": response.get_json()["version"]["import_batch_id"],
        }

    def _wait_for_pipeline(
        self,
        token: str,
        pipeline_run_id: str,
        expected_status: str,
        attempts: int = 40,
    ) -> dict:
        detail_payload = {}
        for _ in range(attempts):
            detail = self.client.get(
                f"/api/v1/admin/pipeline/{pipeline_run_id}",
                headers=self._auth_headers(token),
            )
            self.assertEqual(detail.status_code, 200, detail.get_data(as_text=True))
            detail_payload = detail.get_json()
            current_status = detail_payload["pipeline_run"]["status"]
            if current_status == expected_status:
                return detail_payload
            if current_status == "failed" and expected_status != "failed":
                self.fail(detail_payload)
            time.sleep(0.05)
        self.fail(detail_payload)

    @staticmethod
    def _auth_headers(token: str) -> dict[str, str]:
        return {"Authorization": f"Bearer {token}"}


class FakeIndexWriter:
    def __init__(self, name: str, fail_index: bool = False):
        self.name = name
        self.fail_index = fail_index
        self.indexed_batches: list[str] = []
        self.deleted_batches: list[str] = []
        self.relations_counts: list[int] = []

    def index_chunks(self, chunks, relations=None) -> None:
        if self.fail_index:
            raise RuntimeError(f"{self.name} failed")
        self.indexed_batches.append(chunks[0].import_batch_id)
        self.relations_counts.append(len(relations or []))

    def delete_by_batch(self, import_batch_id: str) -> None:
        self.deleted_batches.append(import_batch_id)


def make_fake_writers() -> list[FakeIndexWriter]:
    return [
        FakeIndexWriter("neo4j"),
        FakeIndexWriter("chroma"),
        FakeIndexWriter("elasticsearch"),
    ]


def complete_publish_metadata() -> dict[str, str]:
    return {
        "title": "Quyet dinh test",
        "document_number": "01/2026/QD-TEST",
        "document_type": "Quyet dinh",
        "issuing_body": "QUOC HOI",
        "issued_date": "2026-01-01",
        "effective_date": "2026-01-01",
        "validity_status": "active",
    }


def make_docx_file() -> io.BytesIO:
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "QUOC HOI"
    table.cell(0, 1).text = "CONG HOA XA HOI CHU NGHIA VIET NAM"
    table.cell(1, 0).text = "So: 01/2026/QD-TEST"
    table.cell(1, 1).text = "Doc lap - Tu do - Hanh phuc"
    document.add_heading("QUYET DINH TEST", level=1)
    document.add_paragraph("Ha Noi, ngay 01 thang 01 nam 2026")
    document.add_paragraph("Dieu 1. Pham vi dieu chinh")
    document.add_paragraph("1. Noi dung khoan mot ve bao hiem.")
    document.add_paragraph("2. Noi dung khoan hai ve hieu luc.")
    document.add_paragraph("Dieu 2. Hieu luc thi hanh")
    document.add_paragraph("Van ban nay co hieu luc ke tu ngay ky.")
    stream = io.BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream


def make_docx_file_with_signature_table() -> io.BytesIO:
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "QUOC HOI"
    table.cell(0, 1).text = "CONG HOA XA HOI CHU NGHIA VIET NAM"
    table.cell(1, 0).text = "So: 116/2025/QH15"
    table.cell(1, 1).text = "Doc lap - Tu do - Hanh phuc"
    document.add_heading("LUAT AN NINH MANG", level=1)
    document.add_paragraph("Ha Noi, ngay 10 thang 12 nam 2025")
    document.add_paragraph("Dieu 1. Pham vi dieu chinh")
    document.add_paragraph("Noi dung ve an ninh mang.")
    document.add_paragraph("Dieu 2. Hieu luc thi hanh")
    document.add_paragraph("Van ban nay co hieu luc tu ngay 01/07/2026.")
    signature_table = document.add_table(rows=2, cols=1)
    signature_table.cell(0, 0).text = "CHU TICH QUOC HOI"
    signature_table.cell(1, 0).text = "Tran Thanh Man"
    stream = io.BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream


def make_docx_file_without_metadata() -> io.BytesIO:
    document = Document()
    document.add_paragraph("No metadata here.")
    document.add_paragraph("This document intentionally has no number or useful title.")
    stream = io.BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream


if __name__ == "__main__":
    unittest.main()
