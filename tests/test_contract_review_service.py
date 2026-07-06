from __future__ import annotations

import io
import json
import os
import tempfile
import unittest
from pathlib import Path

from docx import Document

from backend.config import Config
from backend.models.database import get_connection
from backend.services.auth_service import (
    ROLE_BUSINESS_USER,
    AuthService,
)
from backend.services.chat_agent_service import DocumentStatusRepository
from backend.services.contract_review_service import (
    ContractReviewAgentService,
    DOCUMENT_KIND_CONTRACT,
    DOCUMENT_KIND_LEGAL_DOCUMENT,
    MODULE_AUTHORITY,
    MODULE_EFFECTIVITY,
    detect_document_kind,
    load_enabled_modules,
)


class FakeRetriever:
    def __init__(self):
        self.calls = []

    def search(self, query, top_k, filters, include_expired):
        self.calls.append(
            {
                "query": query,
                "top_k": top_k,
                "filters": filters,
                "include_expired": include_expired,
            }
        )
        return []


class FakeGraphRetriever:
    def enrich(self, hits, top_k):
        return {"related_documents": [], "effectivity_relations": [], "support_score": 0.0}


class ContractReviewHelperTest(unittest.TestCase):
    def test_detect_document_kind_contract_and_legal_document(self) -> None:
        self.assertEqual(
            detect_document_kind("HOP DONG DICH VU\nBen A: Cong ty Bao hiem"),
            DOCUMENT_KIND_CONTRACT,
        )
        self.assertEqual(
            detect_document_kind(
                "CONG HOA XA HOI CHU NGHIA VIET NAM\n"
                "QUOC HOI\n"
                "So: 01/2026/QH15\n"
                "LUAT THU NGHIEM"
            ),
            DOCUMENT_KIND_LEGAL_DOCUMENT,
        )

    def test_module_registry_uses_backend_config_file(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "contract_review_modules.json"
            path.write_text(
                json.dumps({"enabled_modules": [MODULE_AUTHORITY]}),
                encoding="utf-8",
            )

            modules, warnings = load_enabled_modules(path)

        self.assertEqual(modules, [MODULE_AUTHORITY])
        self.assertEqual(warnings, [])


class ContractReviewApiTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.old_cwd = os.getcwd()
        os.chdir(self.temp_dir.name)
        self.db_path = str(Path(self.temp_dir.name) / "app.sqlite3")
        self.modules_path = str(Path(self.temp_dir.name) / "data/rules/contract_review_modules.json")
        self.config = Config(
            app_env="test",
            sqlite_db_path=self.db_path,
            admin_username="admin",
            admin_password="password",
            flask_secret_key="test-secret",
            contract_review_modules_path=self.modules_path,
        )
        from backend.app import create_app

        self.app = create_app(self.config)
        self.client = self.app.test_client()
        self.auth = AuthService(self.db_path)
        self.business_user = self.auth.create_user("business", "password", ROLE_BUSINESS_USER)
        self.other_user = self.auth.create_user("other", "password", ROLE_BUSINESS_USER)
        self.vector = FakeRetriever()
        self.bm25 = FakeRetriever()
        self.review_service = ContractReviewAgentService(
            self.config,
            status_repository=DocumentStatusRepository(self.db_path),
            vector_retriever=self.vector,
            bm25_retriever=self.bm25,
            graph_retriever=FakeGraphRetriever(),
            run_synchronously=True,
        )
        self.app.extensions["contract_review_service"] = self.review_service

    def tearDown(self) -> None:
        os.chdir(self.old_cwd)
        self.temp_dir.cleanup()

    def test_business_user_can_review_and_fetch_completed_job(self) -> None:
        self._insert_published_document("01/2026/QD-TEST", "active")
        token = self._login("business", "password")

        response = self.client.post(
            "/api/v1/contracts/review",
            headers=self._auth_headers(token),
            data={"file": (make_contract_docx(), "contract.docx")},
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 202, response.get_data(as_text=True))
        payload = response.get_json()
        self.assertEqual(payload["status"], "completed")
        self.assertEqual(payload["result"]["document"]["kind"], DOCUMENT_KIND_CONTRACT)
        module_ids = [module["module_id"] for module in payload["result"]["modules"]]
        self.assertEqual(module_ids, [MODULE_AUTHORITY, MODULE_EFFECTIVITY])
        self.assertIn("RULES_MISSING", [item["code"] for item in payload["result"]["warnings"]])
        self.assertTrue(self.vector.calls)
        self.assertTrue(self.bm25.calls)

        fetched = self.client.get(
            f"/api/v1/contracts/review/{payload['job_id']}",
            headers=self._auth_headers(token),
        )
        self.assertEqual(fetched.status_code, 200, fetched.get_data(as_text=True))
        self.assertEqual(fetched.get_json()["result"], payload["result"])

    def test_invalid_docx_is_rejected(self) -> None:
        token = self._login("business", "password")

        response = self.client.post(
            "/api/v1/contracts/review",
            headers=self._auth_headers(token),
            data={"file": (io.BytesIO(b"not a docx"), "contract.docx")},
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.get_json()["error"]["code"], "DOCX_REQUIRED")

    def test_owner_or_admin_can_fetch_job_but_other_business_user_cannot(self) -> None:
        token = self._login("business", "password")
        response = self.client.post(
            "/api/v1/contracts/review",
            headers=self._auth_headers(token),
            data={"file": (make_contract_docx(), "contract.docx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 202, response.get_data(as_text=True))
        job_id = response.get_json()["job_id"]

        other_token = self._login("other", "password")
        forbidden = self.client.get(
            f"/api/v1/contracts/review/{job_id}",
            headers=self._auth_headers(other_token),
        )
        self.assertEqual(forbidden.status_code, 403)

        admin_token = self._login("admin", "password")
        admin_fetch = self.client.get(
            f"/api/v1/contracts/review/{job_id}",
            headers=self._auth_headers(admin_token),
        )
        self.assertEqual(admin_fetch.status_code, 200, admin_fetch.get_data(as_text=True))

    def test_backend_module_config_can_disable_effectivity_module(self) -> None:
        path = Path(self.modules_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(
            json.dumps({"enabled_modules": [MODULE_AUTHORITY]}),
            encoding="utf-8",
        )
        token = self._login("business", "password")

        response = self.client.post(
            "/api/v1/contracts/review",
            headers=self._auth_headers(token),
            data={"file": (make_contract_docx(), "contract.docx")},
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 202, response.get_data(as_text=True))
        modules = response.get_json()["result"]["modules"]
        self.assertEqual([module["module_id"] for module in modules], [MODULE_AUTHORITY])

    def _insert_published_document(self, document_number: str, validity_status: str) -> None:
        with get_connection(self.db_path) as connection:
            connection.execute(
                """
                INSERT INTO document_registry (
                    document_id, document_number, title, source_system, source_url,
                    sector, domain, issuing_body, signer_title, signer_name,
                    document_type, issued_date, effective_date, expiry_date,
                    validity_status, raw_metadata_json, active_version,
                    is_published, is_deleted, created_at, updated_at
                )
                VALUES (?, ?, ?, 'test', '', '', '', '', '', '', '', '', '', '',
                        ?, '{}', 1, 1, 0, '2026-01-01T00:00:00Z', '2026-01-01T00:00:00Z')
                """,
                (
                    f"doc-{document_number}",
                    document_number,
                    f"Document {document_number}",
                    validity_status,
                ),
            )
            connection.commit()

    def _login(self, username: str, password: str) -> str:
        response = self.client.post(
            "/api/v1/auth/login",
            json={"username": username, "password": password},
        )
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        return response.get_json()["access_token"]

    def _auth_headers(self, token: str) -> dict[str, str]:
        return {"Authorization": f"Bearer {token}"}


def make_contract_docx() -> io.BytesIO:
    document = Document()
    document.add_paragraph("HOP DONG DICH VU BAO HIEM")
    document.add_paragraph("Ben A: Cong ty Bao hiem A")
    document.add_paragraph("Ben B: Cong ty Dich vu B")
    document.add_paragraph("Can cu 01/2026/QD-TEST va Dieu 5 cua van ban nay.")
    document.add_paragraph("Dai dien ben A: Giam doc Nguyen Van A")
    document.add_paragraph("Dai dien ben B: Tong giam doc Tran Van B")
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


if __name__ == "__main__":
    unittest.main()
